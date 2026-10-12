"""Giao diện cục bộ của Video Công Thức; worker không truy cập Streamlit."""
from __future__ import annotations

import copy
import hashlib
import logging
import os
import queue
import subprocess
import sys
import tempfile
import threading
import time
from concurrent.futures import ThreadPoolExecutor
from dataclasses import dataclass
from datetime import datetime
from pathlib import Path
from urllib.parse import urlsplit

import streamlit as st
from openpyxl import load_workbook

import applog
import config
import export_excel
import jev_client
import keys
import llm_client
import pipeline
from collectors import cookies, facebook, transcript
from collectors.facebook import page_url_problem
from errors import ServiceError

logger = logging.getLogger("app.ui")


SERVICE_LABELS = {"jev": "Jev", "deepgram": "Deepgram", "llm": "LLM"}
DEFAULT_SETTINGS = {
    "jev": {"model": "", "yes_threshold": 0.5, "min_confidence": 0.6},
    "llm": {"base_url": "", "model": "", "temperature": 0.4},
    "stt": {"provider": "deepgram", "model": "", "language": "vi",
            "keyterms_file": "config/keyterms.txt", "fallback": "local_whisper",
            "whisper_model": "small"},
    "scan": {"video_kinds": ["video", "reel"], "min_views": 10000, "min_ratio_follow": 0, "match_all": True,
             "min_interactions": 200, "delay_min_s": 3, "delay_max_s": 8,
             "page_delay_min_s": 30, "page_delay_max_s": 60,
             "max_videos_per_page": 200, "max_scroll_rounds": 60, "stall_rounds": 4},
    "score_weights": {"ratio_median": 40, "ratio_median_cap": 5, "replicable": 25,
                      "hook_strength": 20, "relevant": 15, "jev_viral": 20},
    "my_page": {"name": "Thế Thảo Pharma", "audience": "Chủ nhà thuốc, dược sĩ"},
}

SERVICE_ROLES = {
    "jev": "Gắn nhãn chủ đề, kiểu hook và cấu trúc cho từng video hoặc bài đăng.",
    "deepgram": "Chuyển lời nói trong video thành chữ.",
    "llm": "Viết phân tích và gợi ý nhân bản. Dùng endpoint tương thích OpenAI.",
}
KEY_HELP = {
    "jev": "Key lấy từ tài khoản Jev. Key chỉ được lưu sau khi kiểm tra thành công.",
    "deepgram": "Key lấy từ trang quản lý API key của Deepgram. Key chỉ được lưu sau khi kiểm tra thành công.",
    "llm": "API key của nhà cung cấp đã chọn. Key chỉ được lưu sau khi kiểm tra thành công.",
}
LLM_PROVIDERS = {
    "OpenRouter": "https://openrouter.ai/api/v1",
    "OpenAI": "https://api.openai.com/v1",
    "Google Gemini": "https://generativelanguage.googleapis.com/v1beta/openai/",
    "Anthropic (Claude)": "https://api.anthropic.com/v1/",
}
LLM_CUSTOM = "Khác (tự nhập)"
PAGE_SETTINGS = "Cài đặt"
PAGE_SCAN = "Chạy quét"
PAGE_RESULTS = "Kết quả"
PAGES = [PAGE_SETTINGS, PAGE_SCAN, PAGE_RESULTS]
XLSX_MIME = "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"


def _redact(text: str, secrets: tuple[str, ...] = ()) -> str:
    for secret in secrets:
        if secret:
            text = text.replace(secret, keys.mask(secret))
    for name in SERVICE_LABELS:
        secret = keys.get_key(name)
        if secret:
            text = text.replace(secret, keys.mask(secret))
    return text


def _error_message(exc: Exception, secrets: tuple[str, ...] = ()) -> str:
    """Câu báo lỗi cho giao diện; đồng thời ghi lỗi vào nhật ký (lỗi lạ kèm traceback)."""
    applog.protect(secrets)
    if isinstance(exc, ServiceError):
        message = _redact(str(exc), secrets)
        logger.warning("%s", message)
        return message
    logger.error("Lỗi không mong đợi: %s", type(exc).__name__, exc_info=exc)
    return ("Không thể hoàn tất thao tác. Kiểm tra kết nối, quyền truy cập và cấu hình rồi thử lại. "
            "Chi tiết trong Nhật ký.")


def _show_error(message: str, exc: BaseException | None = None) -> None:
    """Hiện lỗi vừa xảy ra và ghi vào nhật ký. Chỉ gọi nơi lỗi xảy ra, không ở phần vẽ lại mỗi giây."""
    logger.warning("%s", message, exc_info=exc)
    st.error(message)


def _install_chromium() -> None:
    import browser_setup

    if not config.edge_path():
        browser_setup.ensure_chromium()


@st.cache_resource(show_spinner=False)
def ensure_browser() -> bool:
    """Kiểm tra/tải Chromium một lần; Playwright chạy ngoài thread giao diện."""
    with ThreadPoolExecutor(max_workers=1) as executor:
        try:
            executor.submit(_install_chromium).result()
        except Exception:
            logger.exception("Không chuẩn bị được trình duyệt quét")
            raise
    return True


def _scan_worker(params, settings: dict, events: queue.Queue, stop: threading.Event) -> None:
    def ask_user(kind: str, message: str, options: list[tuple[str, str]]) -> str:
        """Chạy trên thread worker: chờ người dùng chọn; Dừng thì trả về "pause"."""
        question = {"kind": kind, "message": message, "options": options,
                    "answer_event": threading.Event(), "answer": None}
        events.put(("question", question))
        while not question["answer_event"].wait(0.5):
            if stop.is_set():
                return "pause"
        return question["answer"]

    try:
        result = pipeline.run(params, settings, events.put, stop, ask_user=ask_user)
        events.put(("result", result))
    except Exception as exc:
        events.put(("failure", _error_message(exc)))
    finally:
        events.put(("finished", None))


def _login_worker(events: queue.Queue) -> None:
    try:
        success = facebook.login_interactive()
        events.put(("login", success))
    except Exception as exc:
        events.put(("failure", _error_message(exc)))
    finally:
        events.put(("finished", None))


def _init_state() -> None:
    defaults = {"running": False, "logging_in": False, "scan_log": [],
                "scan_errors": [], "progress": None, "scan_result": None, "scan_started": None,
                "key_status": None, "login_message": None, "notice": None, "question": None,
                "svc_models": {}, "page": PAGE_SETTINGS}
    for name, value in defaults.items():
        if name not in st.session_state:
            st.session_state[name] = value


def _busy() -> bool:
    return st.session_state.running or st.session_state.logging_in


def _model_tag(name: str, key: str, base_url: str = "") -> str:
    """Dấu vân tay của key và endpoint: danh sách model chỉ dùng lại khi cả hai không đổi."""
    return hashlib.sha256("\n".join((name, base_url, key)).encode("utf-8")).hexdigest()


def _is_http_url(value: str) -> bool:
    parts = urlsplit(value.strip())
    return parts.scheme in ("http", "https") and bool(parts.netloc)


def _configured(name: str, settings: dict, saved: str | None) -> bool:
    """Đã kết nối khi có key đã lưu và đã chọn đủ model (Deepgram thêm ngôn ngữ, LLM thêm endpoint)."""
    if not saved:
        return False
    if name == "jev":
        return bool(settings["jev"]["model"])
    if name == "deepgram":
        return bool(settings["stt"]["model"] and settings["stt"]["language"])
    return bool(settings["llm"]["base_url"] and settings["llm"]["model"])


def _saved_model(name: str, settings: dict) -> str:
    if name == "jev":
        return settings["jev"]["model"]
    if name == "deepgram":
        return settings["stt"]["model"]
    return settings["llm"]["model"]


def _fetch_models(name: str, candidate: str, base_url: str) -> list[dict]:
    if name == "jev":
        return jev_client.list_models(candidate)
    if name == "deepgram":
        return transcript.list_deepgram_models(candidate)
    return llm_client.list_models(base_url, candidate)


def _connect(name: str, candidate: str, base_url: str) -> None:
    """Tải danh sách model vào session. Không ghi gì vào settings hay kho key."""
    cache = st.session_state.svc_models
    if not candidate:
        cache.pop(name, None)
        _show_error("Vui lòng nhập API key trước khi kết nối.")
        return
    if name == "llm" and not _is_http_url(base_url):
        cache.pop(name, None)
        _show_error("Địa chỉ endpoint phải bắt đầu bằng https:// hoặc http://.")
        return
    try:
        with st.spinner("Đang kết nối..."):
            items = _fetch_models(name, candidate, base_url)
    except Exception as exc:
        cache.pop(name, None)
        st.error(_error_message(exc, (candidate,)))
        return
    if not items:
        cache.pop(name, None)
        _show_error("Dịch vụ không trả về model nào. Kiểm tra lại " + ("key và endpoint." if name == "llm" else "key."))
        return
    cache[name] = {"tag": _model_tag(name, candidate, base_url), "items": items}


def _cached_entry(name: str, candidate: str, base_url: str) -> dict | None:
    entry = st.session_state.svc_models.get(name)
    if entry and candidate and entry["tag"] == _model_tag(name, candidate, base_url):
        return entry
    return None


def _index_of(options: list[str], value: str) -> int:
    return options.index(value) if value in options else 0


def _deepgram_index(items: list[dict], saved: str) -> int:
    names = [item["name"] for item in items]
    if saved in names:
        return names.index(saved)
    # Giá trị cũ như "nova-3" là tên kiến trúc: chọn model đầu tiên thuộc kiến trúc đó.
    for index, item in enumerate(items):
        if saved and item["architecture"] == saved:
            return index
    return 0


def _format_number(value: float) -> str:
    return f"{value:.4f}".rstrip("0").rstrip(".").replace(".", ",")


def _llm_option_label(item: dict) -> str:
    if item["price_in"] is None and item["price_out"] is None:
        price = "giá: nhà cung cấp không công bố qua API"
    else:
        price_in = "không rõ" if item["price_in"] is None else _format_number(item["price_in"])
        price_out = "không rõ" if item["price_out"] is None else _format_number(item["price_out"])
        price = f"giá vào {price_in} · ra {price_out} USD / 1 triệu token"
    parts = [item["id"], price]
    context = item.get("context_length")
    if isinstance(context, (int, float)) and context > 0:
        parts.append(f"ngữ cảnh {int(context):,} token".replace(",", "."))
    return " · ".join(parts)


def _render_jev_model(items: list[dict], settings: dict, tag: str, locked: bool) -> dict:
    names = [item["name"] for item in items]
    model = st.selectbox("Model Jev", names, index=_index_of(names, settings["jev"]["model"]),
                         key=f"svc_model_jev_{tag[:12]}", disabled=locked,
                         help="Model dùng để gắn nhãn video. Mô tả và ngày phát hành hiển thị bên dưới.")
    chosen = next(item for item in items if item["name"] == model)
    details = [chosen["description"], f"phát hành {chosen['release_date']}" if chosen["release_date"] else ""]
    text = " · ".join(part for part in details if part)
    if text:
        st.caption(text)
    return {"model": model}


def _render_deepgram_model(items: list[dict], settings: dict, tag: str, locked: bool) -> dict | None:
    names = [item["name"] for item in items]
    model = st.selectbox("Model Deepgram", names, index=_deepgram_index(items, settings["stt"]["model"]),
                         key=f"svc_model_deepgram_{tag[:12]}", disabled=locked,
                         help="Model chuyển lời nói thành chữ, chỉ gồm model xử lý file ghi sẵn.")
    chosen = next(item for item in items if item["name"] == model)
    if chosen["architecture"]:
        st.caption(f"Kiến trúc: {chosen['architecture']}")
    languages = chosen["languages"]
    if not languages:
        st.error("Model này không có ngôn ngữ nào. Hãy chọn model khác.")
        return None
    codes = [lang["code"] for lang in languages]
    labels = {lang["code"]: f"{lang['label']} ({lang['code']})" for lang in languages}
    saved_language = settings["stt"]["language"]
    if "vi" in codes:
        index = codes.index("vi")
    elif saved_language in codes:
        index = codes.index(saved_language)
    else:
        index = 0
    language = st.selectbox("Ngôn ngữ thoại", codes, index=index, format_func=labels.__getitem__,
                            key=f"svc_lang_{tag[:12]}_{model}", disabled=locked,
                            help="Ngôn ngữ của video. Ưu tiên tiếng Việt (vi) khi model hỗ trợ.")
    return {"model": model, "language": language}


def _render_llm_model(items: list[dict], settings: dict, tag: str, locked: bool, base_url: str) -> dict:
    ids = [item["id"] for item in items]
    by_id = {item["id"]: item for item in items}
    st.caption("Một số nhà cung cấp liệt kê model mà không kiểm tra key. Key chỉ được xác nhận khi bấm Kiểm tra và lưu.")
    model = st.selectbox("Model", ids, index=_index_of(ids, settings["llm"]["model"]),
                         format_func=lambda model_id: _llm_option_label(by_id[model_id]),
                         key=f"svc_model_llm_{tag[:12]}", disabled=locked,
                         help="Gõ tên model vào ô để lọc danh sách. Giá tính bằng USD cho 1 triệu token.")
    return {"model": model, "base_url": base_url}


def _render_selection(name: str, entry: dict, settings: dict, locked: bool, base_url: str) -> dict | None:
    if name == "jev":
        return _render_jev_model(entry["items"], settings, entry["tag"], locked)
    if name == "deepgram":
        return _render_deepgram_model(entry["items"], settings, entry["tag"], locked)
    return _render_llm_model(entry["items"], settings, entry["tag"], locked, base_url)


def _llm_endpoint_input(settings: dict, locked: bool) -> str:
    saved_url = settings["llm"]["base_url"].strip()
    providers = [*LLM_PROVIDERS, LLM_CUSTOM]
    matched = next((label for label, url in LLM_PROVIDERS.items() if url.rstrip("/") == saved_url.rstrip("/")), None)
    default = matched or (LLM_CUSTOM if saved_url else providers[0])
    choice = st.selectbox("Nhà cung cấp", providers, index=providers.index(default), key="svc_llm_provider",
                          disabled=locked,
                          help="Chọn nhà cung cấp có API tương thích OpenAI. Chọn Khác để nhập địa chỉ endpoint riêng.")
    if choice != LLM_CUSTOM:
        st.caption(f"Endpoint: {LLM_PROVIDERS[choice]}")
        return LLM_PROVIDERS[choice]
    custom = st.text_input("Địa chỉ endpoint", value=saved_url if default == LLM_CUSTOM else "",
                           key="svc_llm_custom_url", disabled=locked, placeholder="https://.../v1",
                           help="Đường dẫn API của nhà cung cấp, thường kết thúc bằng /v1.")
    return custom.strip()


def _service_flow(name: str, settings: dict, saved: str | None, locked: bool) -> None:
    st.markdown("**Bước 1 · Kết nối**")
    base_url = _llm_endpoint_input(settings, locked) if name == "llm" else ""
    if saved:
        st.caption(f"Đang lưu key: {keys.mask(saved)}. Để trống ô dưới để kết nối bằng key này.")
    typed = st.text_input("API key", type="password", key=f"svc_key_{name}", disabled=locked,
                          placeholder="Để trống để dùng key đã lưu" if saved else "Dán API key vào đây")
    candidate = typed.strip() or saved or ""
    connected = _cached_entry(name, candidate, base_url) is not None
    if st.button("Kết nối", key=f"svc_connect_{name}", disabled=locked,
                 type="secondary" if connected else "primary",
                 help="Kiểm tra key và tải danh sách model. Bước này chưa lưu gì vào máy."):
        _connect(name, candidate, base_url)
    entry = _cached_entry(name, candidate, base_url)
    if entry is None:
        st.caption("Bấm Kết nối để tải danh sách model. Chưa có gì được lưu cho đến bước 2.")
        return
    st.markdown("**Bước 2 · Chọn model**")
    if st.button("Tải lại danh sách", key=f"svc_refresh_{name}", disabled=locked,
                 help="Tải lại danh sách model từ dịch vụ."):
        _connect(name, candidate, base_url)
        entry = _cached_entry(name, candidate, base_url)
        if entry is None:
            return
    st.caption(f"Đã kết nối · {len(entry['items'])} model khả dụng.")
    selection = _render_selection(name, entry, settings, locked, base_url)
    if selection is None:
        return
    if st.button("Kiểm tra và lưu", type="primary", key=f"svc_save_{name}", disabled=locked,
                 help="Gọi thử model đã chọn một lần. Nếu thành công mới lưu key và model vào máy."):
        _save_service(name, candidate, settings, selection)


def _check_selection(name: str, candidate: str, settings: dict, selection: dict) -> None:
    if name == "jev":
        jev_client.check_key(candidate, selection["model"])
    elif name == "deepgram":
        transcript.check_key(candidate, {**settings["stt"], "model": selection["model"], "language": selection["language"]})
    else:
        llm_client.check_key(selection["base_url"], candidate, selection["model"])


def _apply_selection(settings: dict, name: str, selection: dict) -> None:
    if name == "jev":
        settings["jev"]["model"] = selection["model"]
    elif name == "deepgram":
        settings["stt"].update(model=selection["model"], language=selection["language"])
    else:
        settings["llm"].update(base_url=selection["base_url"], model=selection["model"])


def _save_service(name: str, candidate: str, settings: dict, selection: dict) -> None:
    if not keys.storage_secure():
        _show_error("Không thể lưu API key: kho bảo mật chưa khả dụng. Hãy bật Windows Credential Manager.")
        return
    if not candidate:
        _show_error("Vui lòng nhập API key trước khi kiểm tra.")
        return
    if name == "llm" and not _is_http_url(selection["base_url"]):
        _show_error("Địa chỉ endpoint phải bắt đầu bằng https:// hoặc http://.")
        return
    try:
        with st.spinner("Đang kiểm tra model đã chọn..."):
            _check_selection(name, candidate, settings, selection)
            updated = copy.deepcopy(settings)
            _apply_selection(updated, name, selection)
            config.save_settings(updated)
            keys.set_key(name, candidate)
    except Exception as exc:
        st.error(_error_message(exc, (candidate,)))
        return
    st.session_state.key_status = None
    # Widget state can only be cleared before it is instantiated on the next run.
    st.session_state["clear_key_input"] = f"svc_key_{name}"
    st.session_state.notice = f"Đã kiểm tra và lưu cấu hình {SERVICE_LABELS[name]}."
    st.rerun()


def _reset_configuration() -> None:
    for name in SERVICE_LABELS:
        keys.delete_key(name)
    config.save_settings(copy.deepcopy(DEFAULT_SETTINGS))
    (config.CONFIG_DIR / "channels.txt").write_text("", encoding="utf-8")
    # Mọi khóa session_state của mục Cài đặt có tiền tố svc_ (ô nhập, nút, danh sách model đã tải).
    for key in [key for key in st.session_state.keys() if key.startswith("svc_")]:
        st.session_state.pop(key, None)
    st.session_state.svc_models = {}
    for key in ("channels_editor", "min_followers", "confirm_reset"):
        st.session_state.pop(key, None)
    st.session_state.key_status = None
    st.session_state.notice = "Đã đặt lại API key, cài đặt và danh sách kênh. Dữ liệu và phiên Facebook được giữ nguyên."


def _import_cookie_file(upload) -> None:
    fd, name = tempfile.mkstemp(dir=config.DATA_DIR, prefix="cookie_", suffix=Path(upload.name).suffix)
    os.close(fd)
    path = Path(name)
    imported = False
    try:
        path.write_bytes(upload.getvalue())
        with st.spinner("Đang nạp cookie..."):
            with ThreadPoolExecutor(max_workers=1) as executor:
                imported = executor.submit(facebook.import_cookies, path).result()
    except Exception as exc:
        st.error(_error_message(exc))
        return
    finally:
        path.unlink(missing_ok=True)
    if not imported:
        _show_error("Không tìm thấy cookie đăng nhập Facebook trong file. Hãy xuất cookie khi đã đăng nhập tài khoản phụ rồi thử lại.")
        return
    st.session_state.notice = "Đã nạp cookie và lưu phiên Facebook."
    st.rerun()


def _saved_channels() -> list[str]:
    path = config.CONFIG_DIR / "channels.txt"
    if not path.exists():
        return []
    lines = (raw.strip() for raw in path.read_text(encoding="utf-8").splitlines())
    return [line for line in lines if line and not line.startswith("#") and page_url_problem(line) is None]


def _channel_problems(text: str) -> list[str]:
    problems = []
    for number, raw in enumerate(text.splitlines(), start=1):
        line = raw.strip()
        if line and not line.startswith("#"):
            problem = page_url_problem(line)
            if problem:
                problems.append(f"Dòng {number} ({line[:50]}): {problem}")
    return problems


def _facebook_logged_in() -> bool:
    if not config.COOKIE_FILE.exists():
        return False
    try:
        return cookies.has_login_cookies(cookies.read_netscape(config.COOKIE_FILE))
    except (OSError, ValueError):
        return False


_STEP_STATES = {"done": ":green[✓ Xong]", "current": ":blue[● Đang làm]", "pending": ":gray[○ Chưa làm]"}


@dataclass(frozen=True)
class _Guide:
    """Trạng thái để vẽ thanh 3 bước và việc tiếp theo; được tính lại mỗi lần giao diện chạy."""

    configured: dict[str, bool]
    channel_count: int
    fb_ready: bool
    running: bool
    has_result: bool
    question: bool

    @property
    def missing(self) -> list[str]:
        """Tên các dịch vụ AI chưa kết nối xong."""
        return [SERVICE_LABELS[name] for name in ("jev", "llm") if not self.configured[name]]

    @property
    def setup_done(self) -> bool:
        return self.channel_count > 0 and self.fb_ready


def _journey(guide: _Guide) -> list[tuple[str, str]]:
    """Trạng thái từng bước: done, current hoặc pending."""
    ai = "done" if not guide.missing else "current"
    setup = "done" if guide.setup_done else ("current" if ai == "done" else "pending")
    if guide.running:
        scan = "current"
    elif guide.has_result:
        scan = "done"
    else:
        scan = "current" if ai == "done" and setup == "done" else "pending"
    return [("1 · Dịch vụ AI", ai), ("2 · Kênh & Facebook", setup), ("3 · Chạy quét", scan)]


def _next_step(guide: _Guide) -> tuple[str, str]:
    """Việc nên làm tiếp theo: (hướng dẫn, trang cần mở)."""
    if guide.question:
        return "Quét đang chờ bạn chọn cách xử lý ở tab Chạy quét.", PAGE_SCAN
    if guide.running:
        return "Đang quét. Theo dõi tiến độ ở tab Chạy quét.", PAGE_SCAN
    if guide.missing:
        return f"Kết nối Jev và LLM để quét. Còn thiếu: {', '.join(guide.missing)}. Deepgram chỉ cần cho video.", PAGE_SETTINGS
    if not guide.setup_done:
        todo = [text for text, done in (("thêm kênh cần quét", guide.channel_count > 0),
                                        ("đăng nhập Facebook", guide.fb_ready)) if not done]
        return f"Hoàn tất bước 2: {' và '.join(todo)}.", PAGE_SETTINGS
    if guide.has_result:
        return "Đã có file Excel mới. Xem và tải về ở tab Kết quả.", PAGE_RESULTS
    return "Mọi thứ đã sẵn sàng. Chọn loại quét và ngưỡng rồi bấm Bắt đầu quét.", PAGE_SCAN


def _open_page(page: str) -> None:
    """Callback của nút chuyển trang; chạy trước khi giao diện được dựng lại."""
    st.session_state.page = page


def _render_journey(guide: _Guide) -> None:
    for column, (title, state) in zip(st.columns(3), _journey(guide)):
        with column, st.container(border=True):
            st.markdown(f"**{title}**")
            st.markdown(_STEP_STATES[state])


def _render_next_step(page: str, guide: _Guide) -> None:
    text, target = _next_step(guide)
    with st.container(border=True):
        message, action = st.columns([3, 1], vertical_alignment="center")
        message.markdown(f"**Việc tiếp theo:** {text}")
        if target != page:
            action.button(f"Đi tới {target}", type="primary", width="stretch", on_click=_open_page, args=(target,))


def _render_locked(page: str, guide: _Guide) -> None:
    """Quét cần Jev và LLM; Deepgram chỉ cần cho video."""
    status = "  ·  ".join(f"{'✓' if guide.configured[name] else '✗'} {SERVICE_LABELS[name]}" for name in ("jev", "llm"))
    with st.container(border=True):
        st.markdown(f"**🔒 {page} đang khóa**")
        st.caption(f"Cần kết nối Jev và LLM trước: {status}. Deepgram chỉ cần khi quét video.")
        st.button("Mở Cài đặt", type="primary", on_click=_open_page, args=(PAGE_SETTINGS,))


def _service_card(name: str, settings: dict, locked: bool, *, configured: bool, expanded: bool) -> None:
    label = SERVICE_LABELS[name]
    saved = keys.get_key(name)
    title = f"✓ {label} · model {_saved_model(name, settings)}" if configured else f"○ {label} · chưa cấu hình"
    # Key đổi theo trạng thái để mục vừa xong tự thu gọn và mục kế tiếp tự mở.
    with st.expander(title, expanded=expanded, key=f"svc_card_{name}_{configured}"):
        st.caption(f"{SERVICE_ROLES[name]} {KEY_HELP[name]}")
        if saved and not configured:
            st.warning("Đã có key nhưng còn thiếu model. Làm tiếp hai bước dưới đây.")
        _service_flow(name, settings, saved, locked)


def _channels_section(locked: bool) -> None:
    channels_file = config.CONFIG_DIR / "channels.txt"
    saved_text = channels_file.read_text(encoding="utf-8") if channels_file.exists() else ""
    with st.container(border=True):
        st.markdown("**Kênh cần quét**")
        saved_problems = _channel_problems(saved_text)
        if saved_problems:
            # Danh sách cũ có thể có link nhóm/bài viết: báo ngay, trước khi người dùng bấm Chạy.
            st.warning("Các link sau sẽ bị bỏ qua khi quét:\n\n" + "\n".join(f"- {p}" for p in saved_problems))
        with st.form("channels_form", border=False):
            st.caption("Mỗi dòng là link trang chính của một page (fanpage), ví dụ https://www.facebook.com/ten-page. "
                       "Không hỗ trợ nhóm (group). Dòng bắt đầu bằng # là ghi chú.")
            channels = st.text_area(
                "Danh sách link kênh", height=160, key="channels_editor", disabled=locked, value=saved_text)
            if st.form_submit_button("Lưu danh sách kênh", type="primary", disabled=locked):
                problems = _channel_problems(channels)
                if problems:
                    _show_error("Chưa lưu. Sửa các dòng sau:\n\n" + "\n".join(f"- {p}" for p in problems))
                else:
                    lines = [raw.strip() for raw in channels.splitlines() if raw.strip()]
                    try:
                        channels_file.write_text("\n".join(lines) + "\n", encoding="utf-8")
                    except OSError as exc:
                        _show_error("Không thể lưu danh sách kênh. Kiểm tra quyền ghi thư mục config.", exc)
                    else:
                        count = sum(1 for line in lines if not line.startswith("#"))
                        st.session_state.notice = f"Đã lưu {count} kênh cần quét."
                        st.rerun()


def _facebook_section(fb_ready: bool, locked: bool) -> None:
    with st.container(border=True):
        st.markdown("**Đăng nhập Facebook**")
        st.markdown(":green[✓ Đã đăng nhập]" if fb_ready else ":orange[✗ Chưa đăng nhập]")
        st.caption("Dùng tài khoản Facebook phụ để app quét được video. Cửa sổ trình duyệt sẽ mở để bạn tự đăng nhập; phiên được lưu cho lần quét sau.")
        st.caption("Không chia sẻ thư mục profile/ và file data/fb_cookies.txt: chúng chứa phiên đăng nhập Facebook của bạn.")
        if st.button("Đăng nhập lại" if fb_ready else "Đăng nhập Facebook", type="secondary" if fb_ready else "primary",
                     disabled=locked, icon=":material/login:"):
            events = queue.Queue()
            worker = threading.Thread(target=_login_worker, args=(events,), daemon=True, name="facebook-login")
            st.session_state.login_queue = events
            st.session_state.login_thread = worker
            st.session_state.logging_in = True
            st.session_state.login_message = None
            worker.start()
            st.rerun()
        _login_status()
        with st.expander("Không đăng nhập được? Nạp file cookie"):
            st.caption("Đăng nhập tài khoản Facebook phụ trên trình duyệt, dùng tiện ích Cookie-Editor để xuất cookie của facebook.com ra file JSON rồi tải lên đây. Không chia sẻ file cookie cho người khác.")
            uploaded = st.file_uploader("File cookie (.json hoặc .txt)", type=["json", "txt"], key="cookie_upload", disabled=locked)
            if uploaded is not None and st.button("Nạp cookie", type="primary", disabled=locked):
                _import_cookie_file(uploaded)


def _taxonomy_section(locked: bool) -> None:
    with st.expander("Nhãn phân loại (Jev)"):
        st.caption("Sửa tên và mô tả, thêm dòng cuối bảng hoặc chọn dòng để xóa. Mỗi nhóm cần ít nhất 2 nhãn.")
        st.caption("Thay đổi nhãn khiến Jev phân loại lại nội dung ở lần quét tiếp theo, phát sinh một khoản phí Jev nhỏ.")
        taxonomy = config.load_taxonomy()
        revision = st.session_state.get("taxonomy_revision", 0)
        groups = {}
        with st.form(f"taxonomy_form_{revision}", border=False):
            for group, title in config.TAXONOMY_GROUPS.items():
                st.markdown(f"**{title}**")
                groups[group] = st.data_editor(
                    config.taxonomy_rows(taxonomy, group),
                    column_config={
                        "code": None,
                        "name": st.column_config.TextColumn("Tên hiển thị"),
                        "description": st.column_config.TextColumn("Mô tả cho Jev (càng rõ càng tốt, nên có ví dụ)", width="large"),
                    },
                    column_order=["name", "description"], hide_index=True,
                    num_rows="dynamic", disabled=locked, width="stretch",
                    key=f"taxonomy_{group}_{revision}",
                )
                fixed = [taxonomy["labels"][code] for code in ("khong_co", "khac") if code in taxonomy[group]]
                st.caption("Nhãn luôn được giữ (không chỉnh sửa): " + ", ".join(fixed))
            saved = st.form_submit_button("Lưu nhãn", type="primary", disabled=locked)
        if saved:
            try:
                config.save_taxonomy(groups)
            except (ValueError, OSError) as exc:
                st.error(str(exc))
            else:
                st.session_state.taxonomy_revision = revision + 1
                st.session_state.notice = "Đã lưu nhãn phân loại Jev."
                st.rerun()
        confirmed = st.checkbox("Tôi xác nhận khôi phục nhãn mặc định", disabled=locked, key=f"taxonomy_reset_{revision}")
        if st.button("Khôi phục nhãn mặc định", disabled=locked or not confirmed):
            try:
                config.reset_taxonomy()
            except OSError as exc:
                st.error(str(exc))
            else:
                st.session_state.taxonomy_revision = revision + 1
                st.session_state.notice = "Đã khôi phục nhãn mặc định."
                st.rerun()


def _settings_page(settings: dict, guide: _Guide) -> None:
    locked = _busy()
    st.caption("🔒 Key chỉ được lưu trong Windows Credential Manager, không ghi ra file.")
    if not keys.storage_secure():
        st.warning("Máy này chưa có kho lưu key an toàn. Hãy bật Windows Credential Manager trước khi lưu key.")
    st.subheader("1 · Dịch vụ AI")
    st.caption("Làm lần lượt: nhập key → Kết nối → chọn model → Kiểm tra và lưu.")
    first_open = next((name for name in SERVICE_LABELS if not guide.configured[name]), None)
    for name in SERVICE_LABELS:
        _service_card(name, settings, locked, configured=guide.configured[name], expanded=name == first_open)
    st.subheader("2 · Kênh & Facebook")
    _channels_section(locked)
    _facebook_section(guide.fb_ready, locked)
    _taxonomy_section(locked)
    with st.expander("⚠️ Đặt lại cấu hình"):
        st.caption("Xóa 3 API key, khôi phục settings.json mặc định và xóa danh sách kênh. Không xóa Excel, cache hay phiên Facebook.")
        confirmed = st.checkbox("Tôi xác nhận đặt lại toàn bộ cấu hình trên", key="confirm_reset", disabled=locked)
        if st.button("Đặt lại toàn bộ cấu hình", disabled=locked or not confirmed, icon=":material/restart_alt:"):
            # Defer widget-state deletion until before widgets are built again.
            st.session_state.reset_requested = True
            st.rerun()


def _render_login_status() -> None:
    if st.session_state.logging_in:
        st.info("Đang chờ đăng nhập trong cửa sổ Facebook. Hoàn tất các bước xác minh nếu có.")
    if st.session_state.login_message:
        st.error(st.session_state.login_message)


def _refresh_keys(settings: dict, *, require_deepgram: bool = True) -> dict:
    try:
        with st.spinner("Đang kiểm tra API key..."):
            result = pipeline.check_all_keys(settings, require_deepgram=require_deepgram)
        result = {name: _redact(str(reason)) if reason else None for name, reason in result.items()}
        for name, reason in result.items():
            if reason:
                logger.warning("Key %s: %s", SERVICE_LABELS[name], reason)
    except Exception as exc:
        message = _error_message(exc)
        result = {name: message for name in (SERVICE_LABELS if require_deepgram else ("jev", "llm"))}
    st.session_state.key_status = result
    return result


def _key_status_panel(settings: dict, locked: bool, *, require_deepgram: bool = True) -> None:
    names = tuple(SERVICE_LABELS) if require_deepgram else ("jev", "llm")
    status = st.session_state.key_status
    if status is None or set(status) != set(names):
        status = _refresh_keys(settings, require_deepgram=require_deepgram)
    failed = [SERVICE_LABELS[name] for name in names if status[name]]
    with st.expander("Trạng thái API key", expanded=bool(failed)):
        if failed:
            st.error(f"Chưa thể chạy quét vì lỗi key ở: {', '.join(failed)}. Sửa trong mục Dịch vụ AI rồi kiểm tra lại.")
        for column, name in zip(st.columns(len(names)), names):
            with column:
                reason = status[name]
                st.markdown(f"{':orange[✗]' if reason else ':green[✓]'} **{SERVICE_LABELS[name]}**")
                st.caption(reason or "Key hợp lệ")
        if st.button("Kiểm tra lại", icon=":material/refresh:", disabled=locked):
            _refresh_keys(settings, require_deepgram=require_deepgram)
            st.rerun()


def _scan_options(settings: dict, locked: bool) -> pipeline.RunParams:
    left, right = st.columns(2, gap="medium")
    with left, st.container(border=True):
        st.markdown("**Nội dung nào được quét?**")
        days_choice = st.radio(
            "Nội dung đăng trong", ["7", "14", "30", "Tự nhập"], horizontal=True, disabled=locked,
            format_func=lambda choice: choice if choice == "Tự nhập" else f"{choice} ngày",
            help="Chỉ lấy nội dung đăng trong khoảng này, tính đến hôm nay.")
        days = int(st.number_input("Số ngày", min_value=1, value=7, step=1, disabled=locked)) if days_choice == "Tự nhập" else int(days_choice)
        scan_kind = st.radio("Loại quét", ["Video", "Bài đăng"], horizontal=True, disabled=locked,
                             help="Bài đăng lấy nội dung chữ/ảnh từ bảng tin chính; phân tích caption, không tải audio.")
        formats = {"post"}
        if scan_kind == "Video":
            saved = settings["scan"].get("video_kinds", ["video", "reel"])
            formats = {
                kind for kind, label in (
                    ("video", "Video thường"), ("reel", "Reel"),
                    ("video_live", "Video (Live)"), ("reel_live", "Reel (Live)")
                )
                if st.checkbox(label, value=kind in saved, key=f"scan_kind_{kind}", disabled=locked)
            }
            if not formats:
                st.caption("Chọn ít nhất một loại video để quét.")
        else:
            previous = settings["scan"].get("read_images", False)
            enabled = st.checkbox(
                "Đọc nội dung ảnh trong bài (AI đọc ảnh)", value=previous, disabled=locked,
                help="Tốn thêm khoảng 1.500–2.000 token/ảnh, tối đa 4 ảnh/bài.")
            if enabled and not locked:
                llm = settings["llm"]
                ok = llm_client.vision_cached(llm)
                if ok is None:
                    try:
                        with st.spinner("Đang kiểm tra model LLM đọc ảnh…"):
                            ok = llm_client.check_vision(llm["base_url"], keys.get_key("llm"), llm["model"])
                        llm["vision"] = {"base_url": llm["base_url"], "model": llm["model"], "ok": ok}
                        config.save_settings(settings)
                    except Exception as exc:
                        st.error(_error_message(exc))
                        enabled = False
                if ok is False:
                    st.error(llm_client.vision_error(llm["model"]))
                    enabled = False
                elif ok is True:
                    st.success("Model LLM hỗ trợ đọc ảnh.")
            if not locked and enabled != previous:
                settings["scan"]["read_images"] = enabled
                config.save_settings(settings)
    with right, st.container(border=True):
        if scan_kind == "Bài đăng":
            minimum = int(st.number_input(
                "Tương tác tối thiểu (cảm xúc + bình luận + chia sẻ)", min_value=0,
                value=int(settings["scan"].get("min_interactions", 200)), step=10, key="scan_min_interactions", disabled=locked,
                help="Giữ bài có tổng tương tác đạt hoặc vượt ngưỡng. Không biết số liệu thì không đạt; đặt 0 để không lọc."))
            return pipeline.RunParams(days, formats, min_interactions=minimum)
        minimum = int(st.number_input(
            "Lượt xem tối thiểu", min_value=0, value=int(settings["scan"].get("min_views", 10000)),
            step=1000, key="scan_min_views", disabled=locked,
            help="Giữ video có lượt xem đạt hoặc vượt ngưỡng. Không biết lượt xem thì không đạt; đặt 0 để không lọc."))
        _preset_buttons("scan_min_views", (50_000, 100_000), lambda n: f"{n:,}".replace(",", "."), locked)
        ratio = float(st.number_input(
            "Lượt xem tối thiểu so với follower (×)", min_value=0.0,
            value=float(settings["scan"].get("min_ratio_follow", 0)), step=1.0, format="%g",
            key="scan_min_ratio_follow", disabled=locked,
            help="Ví dụ 5 = lượt xem phải gấp ít nhất 5 lần số follower của page (hay dùng 5× và 10×). "
                 "Không biết follower thì không đạt; đặt 0 để tắt."))
        _preset_buttons("scan_min_ratio_follow", (5.0, 10.0), lambda n: f"{n:g}×", locked)
        match_all = bool(settings["scan"].get("match_all", True))
        if minimum > 0 and ratio > 0:
            both = "Đạt cả hai điều kiện"
            choice = st.radio("Cách kết hợp", [both, "Đạt một trong hai"], index=0 if match_all else 1,
                              horizontal=True, key="scan_match_all", disabled=locked)
            match_all = choice == both
        params = pipeline.RunParams(days, formats, min_views=minimum, min_ratio_follow=ratio, match_all=match_all)
        st.markdown(f"*{_filter_summary(params)}*")
    return params


def _preset_buttons(key: str, values: tuple, label, locked: bool) -> None:
    """Nút chọn nhanh: đặt giá trị vào ô số ngay phía trên."""
    def choose(value) -> None:
        st.session_state[key] = value
    for column, value in zip(st.columns(len(values)), values):
        with column:
            st.button(label(value), key=f"{key}_preset_{value}", on_click=choose, args=(value,),
                      disabled=locked, width="stretch")


def _filter_summary(params: pipeline.RunParams) -> str:
    parts = []
    if params.min_views > 0:
        parts.append(f"≥ {params.min_views:,} lượt xem".replace(",", "."))
    if params.min_ratio_follow > 0:
        parts.append(f"≥ {params.min_ratio_follow:g}× follower")
    if not parts:
        return "Không lọc: giữ mọi video."
    return "Giữ video có " + (" VÀ " if params.match_all else " HOẶC ").join(parts)


def _start_scan(settings: dict, params: pipeline.RunParams) -> None:
    if any(_refresh_keys(settings, require_deepgram=not pipeline.is_post_run(params)).values()):
        st.rerun()  # Panel Trạng thái API key sẽ hiện lỗi sau khi làm mới.
    if not config.load_lines("config/channels.txt"):
        _show_error("Danh sách kênh trống. Vào Cài đặt để nhập các kênh cần quét.")
        return
    if pipeline.is_post_run(params):
        settings["scan"]["min_interactions"] = params.min_interactions
    else:
        settings["scan"].update(min_views=params.min_views, min_ratio_follow=params.min_ratio_follow,
                                match_all=params.match_all, video_kinds=sorted(params.formats))
    config.save_settings(settings)
    events, stop = queue.Queue(), threading.Event()
    worker = threading.Thread(target=_scan_worker, args=(params, copy.deepcopy(settings), events, stop),
                              daemon=True, name="video-scan")
    st.session_state.update(running=True, scan_queue=events, stop_event=stop, scan_thread=worker,
                            scan_started=time.monotonic(), scan_log=[], scan_errors=[], progress=None,
                            scan_result=None, scan_elapsed=None, question=None,
                            scan_item_label="bài đăng" if pipeline.is_post_run(params) else "video")
    worker.start()
    st.rerun()


def _scan_page(settings: dict) -> None:
    _scan_status()
    locked = _busy()
    params = _scan_options(settings, locked)
    _key_status_panel(settings, locked, require_deepgram=not pipeline.is_post_run(params))
    accepted = _terms_gate(settings, locked)
    if st.button("Bắt đầu quét", type="primary", width="stretch",
                 disabled=locked or not params.formats or not accepted, icon=":material/play_arrow:"):
        _start_scan(settings, params)


TERMS_NOTICE = (
    "Công cụ này đọc Facebook qua trình duyệt đã đăng nhập (thu thập tự động). "
    "[Điều khoản Facebook](https://www.facebook.com/terms) hạn chế việc thu thập tự động khi chưa được phép.\n\n"
    "- Tài khoản dùng để quét có thể bị Facebook yêu cầu xác minh (checkpoint) hoặc khóa. "
    "Nên dùng tài khoản phụ, không dùng tài khoản chính.\n"
    "- Chỉ thu thập nội dung bạn được phép truy cập.\n"
    "- Bạn tự chịu trách nhiệm về việc sử dụng công cụ."
)


def _terms_gate(settings: dict, locked: bool) -> bool:
    """Xác nhận một lần trước lần quét đầu tiên; trả về True khi đã xác nhận."""
    accepted_at = settings.get("terms_accepted_at")
    if accepted_at:
        with st.expander("Lưu ý khi dùng Facebook (đã xác nhận)"):
            st.markdown(TERMS_NOTICE)
            st.caption(f"Đã xác nhận lúc {accepted_at}")
        return True
    with st.container(border=True):
        st.markdown("**Lưu ý quan trọng trước khi quét**")
        st.markdown(TERMS_NOTICE)
        ticked = st.checkbox("Tôi đã đọc và tự chịu trách nhiệm", value=False, disabled=locked, key="terms_ack")
    if ticked:
        settings["terms_accepted_at"] = datetime.now().astimezone().isoformat(timespec="seconds")
        config.save_settings(settings)
    return bool(ticked)


def _format_duration(seconds: float) -> str:
    minutes, secs = divmod(int(seconds), 60)
    hours, minutes = divmod(minutes, 60)
    return f"{hours}:{minutes:02d}:{secs:02d}" if hours else f"{minutes:02d}:{secs:02d}"


def _format_size(size: int) -> str:
    if size < 1024 * 1024:
        return f"{size / 1024:.0f} KB"
    return f"{size / (1024 * 1024):.1f} MB".replace(".", ",")


def _exported_video_count(path: Path) -> int | None:
    """Số video trong sheet chi tiết (dòng đầu là tiêu đề). None nếu không đọc được file."""
    try:
        workbook = load_workbook(path, read_only=True)
        try:
            title = export_excel.SHEET_POST_DETAIL if export_excel.SHEET_POST_DETAIL in workbook.sheetnames else export_excel.SHEET_DETAIL
            return max(0, (workbook[title].max_row or 1) - 1)
        finally:
            workbook.close()
    except Exception:  # chỉ là thông tin phụ: file lỗi không được làm hỏng giao diện
        return None


def _download_button(path: Path) -> None:
    try:
        with path.open("rb") as data:
            st.download_button("Tải file Excel", data=data, file_name=path.name, mime=XLSX_MIME,
                               type="primary", icon=":material/download:", key=f"download_{path.name}")
    except OSError:  # vẽ lại mỗi giây trong khung trạng thái quét: không ghi nhật ký ở đây
        st.error("Không đọc được file Excel. Kiểm tra quyền truy cập thư mục output.")


SCAN_STAGES = (("keys", "Kiểm tra key"), ("scan", "Quét kênh"), ("filter", "Lọc nội dung"),
               ("analyze", "Phân tích"), ("export", "Xuất Excel"))
TIMELINE_ROWS = 8
LOG_MAX_RECORDS = 500


def _stage_line(stage: str) -> str:
    names = [name for name, _ in SCAN_STAGES]
    position = len(names) if stage == "done" else names.index(stage) if stage in names else -1
    parts = []
    for index, (_, label) in enumerate(SCAN_STAGES):
        if index < position:
            parts.append(f":green[✓ {label}]")
        elif index == position:
            parts.append(f"**● {label}**")
        else:
            parts.append(f":gray[○ {label}]")
    return "  →  ".join(parts)


def _progress_bar(progress: pipeline.Progress) -> None:
    """Thanh tiến độ của bước hiện tại: số video khi đếm được, nếu không thì số kênh."""
    if progress.steps:
        fraction = progress.step / progress.steps
        item = st.session_state.get("scan_item_label", "video")
        label = (f"Đã phân tích {progress.step}/{progress.steps} {item}" if progress.stage == "analyze"
                 else f"{progress.step}/{progress.steps} {item}")
    elif progress.page_total and progress.stage == "scan":
        fraction = (progress.page_index - 1) / progress.page_total
        label = f"Kênh {progress.page_index}/{progress.page_total}"
    else:
        fraction, label = (1.0, "Xong") if progress.done else (0.0, "")
    st.progress(min(1.0, max(0.0, fraction)), text=label)


def _int_text(value: float | None) -> str:
    return "—" if value is None else f"{round(value):,}".replace(",", ".")


def _usage_panel(usage: dict | None) -> None:
    """Số liệu dịch vụ của lần quét: tốc độ, token, số lần gọi, thời gian chờ LLM; token Jev, phút Deepgram."""
    if not usage:
        return
    speed = pipeline.llm_speed(usage.get("llm_output_tokens", 0), usage.get("llm_seconds", 0.0))
    speed_col, tokens_col, calls_col, wait_col = st.columns(4)
    speed_col.metric("Tốc độ LLM (token/giây)", _int_text(speed),
                     help="Token LLM trả về chia cho thời gian chờ các lần gọi (đo trên máy này).")
    tokens_col.metric("Token LLM", _int_text(usage.get("llm_tokens", 0)),
                      help="Tổng token gửi đi và nhận về; đây là số nhà cung cấp tính phí. "
                           f"Trong đó {_int_text(usage.get('llm_output_tokens', 0))} token trả về.")
    calls_col.metric("Lần gọi LLM", _int_text(usage.get("llm_calls", 0)),
                     help="Đọc ảnh, viết phân tích, viết lại khi Jev loại hết gợi ý, nhận xét tổng hợp.")
    wait_col.metric("Thời gian chờ LLM (giây)", _format_number(round(usage.get("llm_seconds", 0.0), 1)),
                    help="Cộng thời gian chờ của mọi lần gọi; các mục chạy song song nên có thể lớn hơn thời gian quét.")
    cached = usage.get("llm_cached", 0)
    st.caption(f"Model LLM: {usage.get('llm_model') or '—'} · Token Jev: {_int_text(usage.get('jev_tokens', 0))}"
               f" · Deepgram: {_format_number(round(usage.get('deepgram_minutes', 0.0), 2))} phút"
               + (f" · Dùng lại {cached} phân tích LLM đã lưu (không gọi LLM, không tốn token)" if cached else ""))
    if speed is None:
        st.caption("Tốc độ LLM hiện sau lần gọi LLM đầu tiên (khi có video/bài đạt ngưỡng cần phân tích mới).")


def _render_running_scan() -> None:
    progress = st.session_state.progress
    with st.container(border=True):
        elapsed = _format_duration(time.monotonic() - st.session_state.scan_started)
        st.markdown(f"**Đang quét** · đã chạy {elapsed}")
        if progress is not None:
            st.markdown(_stage_line(progress.stage))
            _progress_bar(progress)
            st.markdown(_redact(progress.message))
            if progress.page_url:
                st.caption(_redact(progress.page_url))
            _usage_panel(progress.usage)
        else:
            st.progress(0.0, text="Đang chuẩn bị quét...")
        recent = st.session_state.scan_log[-TIMELINE_ROWS:]
        if recent:
            st.caption("Diễn biến gần nhất")
            st.text("\n".join(recent))
        stopping = st.session_state.stop_event.is_set()
        if st.button("Dừng quét", icon=":material/stop_circle:", key="stop_scan", disabled=stopping):
            st.session_state.stop_event.set()
            st.session_state.question = None
            st.rerun()
        if stopping:
            st.info("Đã yêu cầu dừng. Đang hoàn tất bước hiện tại và lưu kết quả đã xử lý.")
        question = st.session_state.question
        if question is not None:
            st.warning(question["message"])
            for column, (option_id, label) in zip(st.columns(len(question["options"])), question["options"]):
                with column:
                    if st.button(label, key=f"answer_{question['kind']}_{option_id}"):
                        question["answer"] = option_id
                        question["answer_event"].set()
                        st.session_state.question = None
                        st.rerun()


def _render_scan_result(path: Path) -> None:
    count = _exported_video_count(path)
    with st.container(border=True):
        item = st.session_state.get("scan_item_label", "video")
        st.markdown(f"**✅ Quét xong · {count} {item}**" if count is not None else "**✅ Quét xong**")
        elapsed = st.session_state.get("scan_elapsed")
        st.caption(f"File: {path.name}" + (f" · tổng thời gian {_format_duration(elapsed)}" if elapsed else ""))
        progress = st.session_state.get("progress")
        _usage_panel(progress.usage if progress is not None else None)
        _download_button(path)


def _render_scan_status() -> None:
    if st.session_state.running:
        _render_running_scan()
    elif st.session_state.scan_result:
        _render_scan_result(Path(st.session_state.scan_result))
    elif st.session_state.scan_log and not st.session_state.scan_errors:
        st.info("Đã kết thúc quét; không có file Excel mới.")
    for message in dict.fromkeys(st.session_state.scan_errors):
        st.error(message)
    if st.session_state.scan_errors:
        st.caption("Chi tiết kỹ thuật của từng lỗi nằm trong Nhật ký (nút ở góc trên bên phải).")
    if st.session_state.scan_log:
        with st.expander(f"Toàn bộ diễn biến lần quét ({len(st.session_state.scan_log)} bước)"):
            st.text("\n".join(st.session_state.scan_log))


def _drain_scan_queue() -> bool:
    """Nhận sự kiện từ luồng quét. Trả về True khi luồng đã kết thúc."""
    finished = False
    if "scan_queue" in st.session_state:
        while True:
            try:
                event = st.session_state.scan_queue.get_nowait()
            except queue.Empty:
                break
            if isinstance(event, tuple):
                kind, value = event
                if kind == "result":
                    st.session_state.scan_result = str(value) if value else None
                elif kind == "failure":
                    st.session_state.scan_errors.append(value)
                elif kind == "question":
                    st.session_state.question = value
                elif kind == "finished":
                    st.session_state.running = False
                    st.session_state.scan_elapsed = time.monotonic() - st.session_state.scan_started
                    st.session_state.question = None
                    finished = True
            else:
                stamp = datetime.now().strftime("%H:%M:%S")
                if event.message and not event.transient:  # đếm ngược chỉ hiện ở khung tiến độ
                    marker = "⚠ " if event.error else ""
                    st.session_state.scan_log.append(f"{stamp}  {marker}{_redact(event.message)}")
                if event.error:
                    st.session_state.scan_errors.append(_redact(event.error))
                else:  # lỗi không đổi bước đang chạy và thanh tiến độ
                    st.session_state.progress = event
    return finished


def _drain_login_queue() -> bool:
    """Nhận sự kiện đăng nhập Facebook. Trả về True khi cửa sổ đăng nhập đã đóng."""
    finished = False
    if st.session_state.logging_in:
        while True:
            try:
                kind, value = st.session_state.login_queue.get_nowait()
            except queue.Empty:
                break
            if kind == "login":
                if value:
                    st.session_state.login_message = None
                    st.session_state.notice = "Đã đăng nhập Facebook và lưu phiên."
                else:
                    st.session_state.login_message = "Chưa đăng nhập Facebook thành công. Vui lòng thử lại."
                    logger.warning("%s", st.session_state.login_message)
            elif kind == "failure":
                st.session_state.login_message = value
            elif kind == "finished":
                st.session_state.logging_in = False
                finished = True
    return finished


def _poll_workers() -> None:
    """Nhận sự kiện của luồng quét và đăng nhập ở mọi trang; làm mới giao diện khi một trong hai vừa xong."""
    scan_done = _drain_scan_queue()
    login_done = _drain_login_queue()
    if scan_done or login_done:
        st.rerun()


def _results_page() -> None:
    files = sorted(config.OUTPUT_DIR.glob("*.xlsx"), key=lambda path: path.stat().st_mtime, reverse=True)
    title, folder = st.columns([3, 1], vertical_alignment="center")
    title.markdown("**File Excel đã xuất** · mới nhất trước")
    if folder.button("Mở thư mục", icon=":material/folder_open:", width="stretch"):
        try:
            os.startfile(config.OUTPUT_DIR)
        except (OSError, AttributeError) as exc:
            _show_error("Không thể mở thư mục kết quả trên máy này.", exc)
    if not files:
        st.info("Chưa có file Excel nào. Kết quả sẽ hiện ở đây sau mỗi lần quét.")
        return
    for path in files:
        stat = path.stat()
        with st.container(border=True):
            info, action = st.columns([3, 1], vertical_alignment="center")
            info.markdown(f"**{path.name}**")
            info.caption(f"{datetime.fromtimestamp(stat.st_mtime):%d/%m/%Y %H:%M} · {_format_size(stat.st_size)}")
            with action:
                _download_button(path)


@st.dialog("Nhật ký ứng dụng", width="large")
def _log_dialog() -> None:
    """Thông báo và lỗi đã ghi vào data/app.log, mới nhất ở trên; API key luôn được che."""
    files = {"Lần mở này": applog.LOG_FILE}
    if applog.PREVIOUS_LOG_FILE.exists():
        files["Lần mở trước"] = applog.PREVIOUS_LOG_FILE
    session, scope = st.columns([3, 2], vertical_alignment="center")
    choice = (session.radio("Phiên", list(files), horizontal=True, label_visibility="collapsed")
              if len(files) > 1 else next(iter(files)))
    problems_only = scope.toggle("Chỉ lỗi và cảnh báo", value=True)
    path = files[choice]
    records = applog.read_records(problems_only, path)
    st.caption(f"{path} · mới nhất ở trên")
    if records:
        shown = records[-LOG_MAX_RECORDS:][::-1]
        st.code(applog.redact("\n".join(shown)), language=None, wrap_lines=True, height=420)
    elif problems_only:
        st.success("Không có lỗi hay cảnh báo nào.")
    else:
        st.info("Nhật ký đang trống.")
    if st.button("Mở file nhật ký", icon=":material/description:"):
        try:
            os.startfile(path)
        except (OSError, AttributeError) as exc:
            _show_error("Không mở được file nhật ký trên máy này.", exc)


# Fragments refresh only their own part every second; workers communicate solely through queues.
if callable(getattr(st, "fragment", None)):
    _poll = st.fragment(run_every=1)(_poll_workers)
    _scan_status = st.fragment(run_every=1)(_render_scan_status)
    _login_status = st.fragment(run_every=1)(_render_login_status)
else:
    _poll = _poll_workers
    _scan_status = _render_scan_status
    _login_status = _render_login_status


def main() -> None:
    st.set_page_config(page_title="Video Công Thức", page_icon="🎬", layout="wide")
    applog.setup()  # chạy bằng launcher thì đã bật sẵn; `streamlit run app.py` thì bật ở đây
    config.ensure_dirs()
    _init_state()
    # Cửa sổ ứng dụng riêng không có thanh công cụ Streamlit: bỏ khoảng trống mặc định ở đầu trang.
    st.markdown("<style>.block-container{padding-top:1.5rem}</style>", unsafe_allow_html=True)
    title, log_button = st.columns([5, 1], vertical_alignment="center")
    title.markdown("### 🎬 Video Công Thức")
    if log_button.button("Nhật ký", icon=":material/receipt_long:", width="stretch",
                         help="Mọi thông báo và lỗi của ứng dụng (data/app.log)."):
        _log_dialog()
    st.caption("Tìm video hoặc bài đăng nổi bật trên các kênh Facebook để học công thức nội dung, rồi xuất kết quả ra Excel.")
    try:
        if st.session_state.pop("reset_requested", False):
            _reset_configuration()
        clear_input = st.session_state.pop("clear_key_input", None)
        if clear_input:
            st.session_state[clear_input] = ""
        settings = config.load_settings()
        configured = {name: _configured(name, settings, keys.get_key(name)) for name in SERVICE_LABELS}
    except Exception as exc:
        st.error(_error_message(exc))
        return
    if st.session_state.notice:
        st.toast(st.session_state.notice, icon="✅")
        st.session_state.notice = None
    try:
        with st.spinner("Đang tải trình duyệt (~150 MB)..."):
            ensure_browser()
    except Exception:
        st.error("Không thể tải trình duyệt Chromium. Kiểm tra mạng, dung lượng ổ đĩa và mở lại ứng dụng.")
        return
    guide = _Guide(configured=configured, channel_count=len(_saved_channels()), fb_ready=_facebook_logged_in(),
                   running=st.session_state.running, has_result=bool(st.session_state.scan_result),
                   question=st.session_state.question is not None)
    _render_journey(guide)
    page = st.segmented_control("Trang", PAGES, key="page", required=True, width="stretch",
                                label_visibility="collapsed") or PAGE_SETTINGS
    if guide.missing and page == PAGE_SCAN:
        _render_locked(page, guide)
    else:
        _render_next_step(page, guide)
        if page == PAGE_SETTINGS:
            _settings_page(settings, guide)
        elif page == PAGE_SCAN:
            _scan_page(settings)
        else:
            _results_page()
    _poll()
    if not callable(getattr(st, "fragment", None)) and _busy():
        time.sleep(1)
        st.rerun()


if __name__ == "__main__":
    main()
