"""Excel export (PLAN §8): sheet "Chi tiết video", "Tổng hợp công thức" and "Nhật ký"."""
from __future__ import annotations

import json
import logging
import math
from dataclasses import dataclass
from datetime import datetime, timedelta, timezone
from pathlib import Path
from typing import Any, Callable, Collection

from openpyxl import Workbook
from openpyxl.cell.cell import ILLEGAL_CHARACTERS_RE
from openpyxl.chart import BarChart, Reference
from openpyxl.formatting.rule import ColorScaleRule, DataBarRule, FormulaRule
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.worksheet import Worksheet

import config
from models import KIND_LABELS, Video, video_kind
from normalize import to_vn_datetime

log = logging.getLogger(__name__)

SHEET_DETAIL = "Chi tiết video"
SHEET_POST_DETAIL = "Chi tiết bài đăng"
SHEET_SUMMARY = "Tổng hợp công thức"
SHEET_LOG = "Nhật ký"

VN_TZ = timezone(timedelta(hours=7))
UNCERTAIN_SUFFIX = " (chưa chắc)"
NO_LABEL = "Chưa gắn nhãn"
MAX_CELL_CHARS = 32767  # Excel's per-cell limit

_UNCERTAIN_NAMES = {
    "topic": "Nhóm chủ đề",
    "hook_type": "Kiểu hook",
    "hook_strength": "Độ mạnh hook",
    "structure": "Cấu trúc",
    "cta": "CTA",
    "relevant": "Liên quan ngành",
    "replicable": "Làm lại được",
    "cam_xuc_manh": "Cảm xúc mạnh",
    "gia_tri_thuc_tien": "Giá trị thực tiễn",
    "ket_thuc_tron_y": "Kết thúc trọn ý",
    "kha_nang_chia_se": "Khả năng chia sẻ",
}
_PARAM_LABELS = {
    "days": "Số ngày quét",
    "formats": "Định dạng",
    "min_views": "Lượt xem tối thiểu",
    "min_interactions": "Tương tác tối thiểu",
    "min_ratio_follow": "Lượt xem tối thiểu so với follower (×)",
    "match_all": "Cách kết hợp điều kiện",
}

FMT_INT = "#,##0"
FMT_DEC1 = "#,##0.0"
FMT_DEC2 = "#,##0.00"
FMT_PCT0 = "0%"
FMT_PCT1 = "0.0%"
FMT_PCT2 = "0.00%"
FMT_DATETIME = "dd/mm/yyyy hh:mm"

RISK_FILL = PatternFill("solid", fgColor="FCE4E4")
HEADER_FILL = PatternFill("solid", fgColor="234E63")
HEADER_FONT = Font(name="Calibri", bold=True, color="FFFFFF", size=11)
TITLE_FONT = Font(name="Calibri", bold=True, color="234E63", size=14)
BODY_FONT = Font(name="Calibri", color="243746", size=11)
ZEBRA_FILL = PatternFill("solid", fgColor="F3F7FA")
ROW_BORDER = Border(bottom=Side(style="hair", color="DCE5EB"))
TOP = Alignment(vertical="top")
TOP_WRAP = Alignment(vertical="top", wrap_text=True)
HEADER_ALIGN = Alignment(vertical="center", wrap_text=True)

LONG_WIDTH = 60


@dataclass(frozen=True)
class _Column:
    header: str
    get: Callable[[Video, int], Any]  # (video, 1-based STT) -> cell value
    fmt: str | None = None
    width: int = 14
    long: bool = False  # wrap text, width 60
    link: bool = False  # make the cell a hyperlink


def _jev(v: Video, key: str) -> Any:
    return (v.jev or {}).get(key)


def _llm(v: Video, key: str) -> Any:
    return (v.llm or {}).get(key)


def _text(value: Any) -> str:
    return "" if value is None else str(value).strip()


def _as_list(value: Any) -> list:
    if isinstance(value, str):
        return [value] if value.strip() else []
    if isinstance(value, (list, tuple)):
        return list(value)
    return []


def _naive_vn(ts: int | float | None) -> datetime | None:
    if ts is None:
        return None
    return to_vn_datetime(int(ts)).replace(tzinfo=None)  # Excel dates carry no tzinfo


def _risk_flags(v: Video) -> list[str]:
    return [text for text in (_text(f) for f in _as_list(_jev(v, "risk_flags"))) if text]


def _suggestion(v: Video, index: int) -> str:
    items = _as_list(_llm(v, "goi_y_nhan_ban"))
    if index >= len(items):
        return ""
    item = items[index]
    if not isinstance(item, dict):
        return _text(item)
    lines = []
    hook = _text(item.get("hook_moi"))
    idea = _text(item.get("y_tuong"))
    if hook:
        lines.append(f"Hook: {hook}")
    if idea:
        lines.append(f"Ý tưởng: {idea}")
    return "\n".join(lines)


def _notes(v: Video) -> str:
    parts = [text for text in (_text(n) for n in (v.notes or [])) if text]
    labels = [_uncertain_label(x) for x in _as_list(_jev(v, "uncertain"))]
    labels = [label for label in labels if label]
    if labels:
        parts.append("Nhãn chưa chắc: " + ", ".join(labels))
    if v.transcript_status == "failed":
        parts.append("Chưa lấy được thoại")
    raw = _text(_llm(v, "_raw"))
    if raw:
        parts.append("LLM trả về không đọc được JSON, phản hồi thô: " + raw)
    return "\n".join(parts)


def _uncertain_label(key: Any) -> str:
    text = _text(key)
    return _UNCERTAIN_NAMES.get(text, text)


def _score_order(v: Video) -> tuple[bool, float]:
    # Highest score first; videos without a score go last.
    return (v.score is None, -(v.score or 0.0))


COLUMNS: tuple[_Column, ...] = (
    # Định danh
    _Column("STT", lambda v, n: n, FMT_INT, width=7),
    _Column("Kênh", lambda v, _: v.page_name or v.page_url, width=24),
    _Column("Follower kênh", lambda v, _: v.page_followers, FMT_INT),
    _Column("Ngày đăng", lambda v, _: _naive_vn(v.posted_at), FMT_DATETIME, width=17),
    _Column("Định dạng", lambda v, _: KIND_LABELS.get(video_kind(v.format, v.is_live), v.format), width=14),
    _Column("Thời lượng (giây)", lambda v, _: v.duration_s, FMT_INT, width=13),
    _Column("Link bài", lambda v, _: v.url, width=40, link=True),
    # Chỉ số
    _Column("View", lambda v, _: v.views, FMT_INT),
    _Column("Cảm xúc", lambda v, _: v.reactions, FMT_INT),
    _Column("Bình luận", lambda v, _: v.comments, FMT_INT),
    _Column("Chia sẻ", lambda v, _: v.shares, FMT_INT),
    _Column("Tổng tương tác", lambda v, _: v.interactions, FMT_INT),
    _Column("View/Follower", lambda v, _: v.ratio_follow, FMT_PCT1, width=16),
    _Column("Hệ số so trung vị kênh", lambda v, _: v.ratio_median, FMT_PCT1, width=18),
    _Column("Tỷ lệ tương tác", lambda v, _: v.engagement, FMT_PCT1, width=16),
    _Column("Tỷ lệ share", lambda v, _: v.share_rate, FMT_PCT2, width=14),
    # Nội dung
    _Column("Chủ đề bài", lambda v, _: _llm(v, "chu_de_bai"), width=28),
    _Column("Tóm tắt", lambda v, _: _llm(v, "tom_tat"), long=True),
    _Column("Caption", lambda v, _: v.caption, long=True),
    _Column("Nội dung thoại", lambda v, _: v.transcript, long=True),
    _Column("Hook gốc", lambda v, _: _llm(v, "hook_goc") or v.hook, long=True),
    _Column("Góc nhìn", lambda v, _: _llm(v, "goc_nhin"), long=True),
    _Column("Mẫu hook điền chỗ trống", lambda v, _: _llm(v, "mau_hook_dien_cho_trong"), long=True),
    _Column("Vị trí CTA", lambda v, _: _llm(v, "vi_tri_cta"), long=True),
    # Công thức
    _Column("Nhóm chủ đề", lambda v, _: _jev(v, "topic"), width=24),
    _Column("Kiểu hook", lambda v, _: _jev(v, "hook_type"), width=22),
    _Column("Độ mạnh hook", lambda v, _: _jev(v, "hook_strength"), FMT_INT, width=12),
    _Column("Cấu trúc", lambda v, _: _jev(v, "structure"), width=22),
    _Column("CTA", lambda v, _: _jev(v, "cta"), width=22),
    # Đánh giá
    _Column("Liên quan ngành (%)", lambda v, _: _jev(v, "relevant_p"), FMT_PCT0, width=14),
    _Column("Làm lại được (%)", lambda v, _: _jev(v, "replicable_p"), FMT_PCT0, width=14),
    _Column("Điểm công thức (0–100)", lambda v, _: v.score, FMT_DEC1, width=14),
    _Column("Điểm Jev (0–100)", lambda v, _: v.jev["jev_viral"] * 100 if v.jev.get("jev_viral") is not None else None, FMT_DEC1, width=14),
    _Column("Cờ rủi ro", lambda v, _: _risk_flags(v), width=26),
    # Nhân bản
    _Column("Lý do viral", lambda v, _: _llm(v, "ly_do_viral"), long=True),
    _Column("Gợi ý 1", lambda v, _: _suggestion(v, 0), long=True),
    _Column("Gợi ý 2", lambda v, _: _suggestion(v, 1), long=True),
    _Column("Gợi ý 3", lambda v, _: _suggestion(v, 2), long=True),
    # Khác
    _Column("Ghi chú", lambda v, _: _notes(v), long=True),
)


def export(
    videos: list[Video],
    summary: dict,
    run_log: dict,
    out_dir: str | Path = config.OUTPUT_DIR,
    now: datetime | None = None,
) -> Path:
    """Write `cong-thuc_YYYY-MM-DD_HHmm.xlsx` (Vietnam time) and return its path.

    If a file with the same name already exists, `_2`, `_3`, ... is appended so earlier
    exports are never overwritten.
    """
    moment = _moment(now)
    folder = Path(out_dir)
    folder.mkdir(parents=True, exist_ok=True)
    path = _unique_path(folder, moment)

    wb = Workbook()
    detail = wb.active
    post_run = set((run_log.get("params") or {}).get("formats") or []) == {"post"}
    detail.title = SHEET_POST_DETAIL if post_run else SHEET_DETAIL
    _write_details(detail, videos, post_run=post_run)
    _write_summary(wb.create_sheet(SHEET_SUMMARY), summary or {}, post_run=post_run, moment=moment)
    _write_run_log(wb.create_sheet(SHEET_LOG), run_log or {})
    for ws, colour in zip(wb.worksheets, ("287D8E", "55A58B", "8295A3")):
        ws.sheet_properties.tabColor = colour
        ws.sheet_view.showGridLines = False
        ws.page_setup.orientation = "landscape"
        ws.page_setup.paperSize = ws.PAPERSIZE_A3 if ws is detail else ws.PAPERSIZE_A4
        ws.page_setup.fitToWidth = 1
        ws.page_setup.fitToHeight = 0
        ws.sheet_properties.pageSetUpPr.fitToPage = True
        ws.print_options.horizontalCentered = True
        if not ws.print_area:
            ws.print_area = ws.dimensions
    wb.active = 0
    wb.save(path)
    log.info("Đã xuất Excel: %s (%d %s)", path.name, len(videos), "bài đăng" if post_run else "video")
    return path


def build_summary_tables(videos: list[Video]) -> dict:
    """Group by Jev topic, hook and structure, ranked by average formula score.

    The " (chưa chắc)" suffix is stripped for grouping. The commentary is left "" for the LLM.
    """
    return {
        "by_topic": _group(videos, lambda v: _label(v, "topic")),
        "by_hook": _group(videos, lambda v: _label(v, "hook_type")),
        "by_structure": _group(videos, lambda v: _label(v, "structure")),
        "commentary": "",
    }


def _label(v: Video, key: str) -> str:
    text = _text(_jev(v, key)).removesuffix(UNCERTAIN_SUFFIX).strip()
    return text or NO_LABEL


def _group(videos: list[Video], key_of: Callable[[Video], str]) -> list[dict]:
    buckets: dict[str, list[Video]] = {}
    for v in videos:
        buckets.setdefault(key_of(v), []).append(v)
    rows = [
        {
            "label": label,
            "count": len(items),
            "avg_views": _mean([v.views for v in items]),
            "avg_score": _mean([v.score for v in items]),
        }
        for label, items in buckets.items()
    ]
    rows.sort(key=lambda r: (r["avg_score"] is None, -(r["avg_score"] or 0), -r["count"], r["label"]))
    return rows


def _mean(values: list[float | int | None]) -> float | None:
    nums = [x for x in values if x is not None]
    return sum(nums) / len(nums) if nums else None


def _moment(now: datetime | None) -> datetime:
    if now is None:
        return datetime.now(VN_TZ)
    if now.tzinfo is not None:
        return now.astimezone(VN_TZ)
    return now


def _unique_path(folder: Path, moment: datetime) -> Path:
    stem = f"cong-thuc_{moment:%Y-%m-%d_%H%M}"
    path = folder / f"{stem}.xlsx"
    n = 2
    while path.exists():
        path = folder / f"{stem}_{n}.xlsx"
        n += 1
    return path


def _cell_value(value: Any) -> Any:
    """Make a value safe for openpyxl: no illegal control chars, no NaN, lists joined as text."""
    if value is None or isinstance(value, datetime | int):
        return value
    if isinstance(value, float):
        return None if math.isnan(value) else value
    if isinstance(value, (list, tuple)):
        value = ", ".join(t for t in (_text(x) for x in value) if t)
    elif isinstance(value, (set, frozenset)):
        value = ", ".join(sorted(t for t in (_text(x) for x in value) if t))
    elif isinstance(value, dict):
        value = json.dumps(value, ensure_ascii=False)
    text = ILLEGAL_CHARACTERS_RE.sub("", str(value)).strip()
    return text[:MAX_CELL_CHARS] or None


def _put(
    ws: Worksheet,
    row: int,
    col: int,
    value: Any,
    fmt: str | None = None,
    align: Alignment | None = None,
):
    cell = ws.cell(row=row, column=col)
    cell.value = _cell_value(value)
    if isinstance(cell.value, str) and cell.value.startswith("="):
        cell.data_type = "s"  # keep as literal text, never a formula
    if fmt:
        cell.number_format = fmt
    if align is not None:
        cell.alignment = align
    return cell


# Bài đăng không có lượt xem, thời lượng hay thoại: ẩn các cột luôn trống này (vẫn giữ để bố cục hai loại file giống nhau).
_VIDEO_ONLY_HEADERS = frozenset({
    "Thời lượng (giây)", "View", "View/Follower", "Tỷ lệ tương tác", "Tỷ lệ share", "Nội dung thoại",
})


def _write_details(ws: Worksheet, videos: list[Video], *, post_run: bool = False) -> None:
    groups = (
        (1, 7, "Thông tin bài", "DCEAF3"),
        (8, 16, "Chỉ số", "DDF0E8"),
        (17, 24, "Nội dung", "EEE7F5"),
        (25, 34, "Công thức (Jev)", "FFF0D6"),
        (35, 39, "Gợi ý làm lại", "F8E3DA"),
    )
    for first, last, title, colour in groups:
        ws.merge_cells(start_row=1, start_column=first, end_row=1, end_column=last)
        for c in range(first, last + 1):
            ws.cell(1, c).fill = PatternFill("solid", fgColor=colour)
        cell = _put(ws, 1, first, title, align=HEADER_ALIGN)
        cell.font = TITLE_FONT
    has_images = post_run and any(v.transcript_status == "images" and v.transcript for v in videos)
    for c, col in enumerate(COLUMNS, start=1):
        title = "Nội dung ảnh" if has_images and col.header == "Nội dung thoại" else col.header
        header = _put(ws, 2, c, title, align=HEADER_ALIGN)
        header.font = HEADER_FONT
        header.fill = HEADER_FILL
        width = 16 if col.link else LONG_WIDTH if col.long else col.width
        ws.column_dimensions[get_column_letter(c)].width = width
        if post_run and col.header in _VIDEO_ONLY_HEADERS and not (has_images and col.header == "Nội dung thoại"):
            ws.column_dimensions[get_column_letter(c)].hidden = True
    ws.row_dimensions[1].height = 28
    ws.row_dimensions[2].height = 44

    ordered = sorted(videos, key=_score_order)
    for r, video in enumerate(ordered, start=3):
        height = 30.0
        for c, col in enumerate(COLUMNS, start=1):
            wrap = col.long or col.header in ("Cờ rủi ro", "Chủ đề bài")
            fmt = FMT_DEC2 if col.header == "Hệ số so trung vị kênh" else col.fmt
            cell = _put(ws, r, c, col.get(video, r - 2), fmt, TOP_WRAP if wrap else TOP)
            cell.font = BODY_FONT
            cell.border = ROW_BORDER
            if r % 2:
                cell.fill = ZEBRA_FILL
            if wrap and isinstance(cell.value, str):
                height = max(height, _text_height(cell.value, chars_per_line=LONG_WIDTH if col.long else col.width))
            if col.link and isinstance(cell.value, str) and cell.value.startswith("http"):
                cell.hyperlink = cell.value
                cell.hyperlink.display = "Mở bài ↗"
                # Retain the source URL as the cell value, but show a compact link in Excel.
                cell.number_format = ';;;"Mở bài ↗"'
                cell.font = Font(name="Calibri", color="0563C1", underline="single", size=11)
            if col.header == "Cờ rủi ro" and cell.value:
                cell.fill = RISK_FILL
        ws.row_dimensions[r].height = min(90, height)

    ws.freeze_panes = "C3"
    ws.auto_filter.ref = f"A2:{get_column_letter(len(COLUMNS))}{len(ordered) + 2}"
    ws.print_title_rows = "1:2"
    ws.sheet_view.zoomScale = 85
    if ordered:
        end = len(ordered) + 2
        for c, col in enumerate(COLUMNS, start=1):
            letter = get_column_letter(c)
            region = f"{letter}3:{letter}{end}"
            if col.header.startswith(("Điểm công thức", "Điểm Jev")):
                ws.conditional_formatting.add(region, DataBarRule(
                    start_type="num", start_value=0, end_type="num", end_value=100,
                    color="65B5A2", showValue=True,
                ))
            elif col.header.startswith(("Liên quan ngành", "Làm lại được")):
                ws.conditional_formatting.add(region, ColorScaleRule(
                    start_type="num", start_value=0, start_color="F5D8D5",
                    mid_type="num", mid_value=0.5, mid_color="FFF3CE",
                    end_type="num", end_value=1, end_color="CFEBDD",
                ))
            elif col.header == "Cờ rủi ro":
                ws.conditional_formatting.add(region, FormulaRule(
                    formula=[f'LEN({letter}3)>0'], fill=RISK_FILL,
                ))


def _write_table(
    ws: Worksheet,
    start: int,
    title: str,
    headers: list[str],
    rows: list[list[Any]],
    fmts: list[str | None] | None = None,
    wrap: Collection[int] = (),
) -> int:
    """Write a titled table starting at row `start`. Returns the first free row after a blank line."""
    _put(ws, start, 1, title).font = TITLE_FONT
    ws.merge_cells(start_row=start, start_column=1, end_row=start, end_column=len(headers))
    ws.row_dimensions[start].height = 28
    header_row = start + 1
    for c, h in enumerate(headers, start=1):
        cell = _put(ws, header_row, c, h, align=HEADER_ALIGN)
        cell.font = HEADER_FONT
        cell.fill = HEADER_FILL
    ws.row_dimensions[header_row].height = 34
    for i, values in enumerate(rows, start=1):
        r = header_row + i
        for c, value in enumerate(values, start=1):
            fmt = fmts[c - 1] if fmts else None
            cell = _put(ws, r, c, value, fmt, TOP_WRAP if c in wrap else TOP)
            cell.font = BODY_FONT
            cell.border = ROW_BORDER
            if i % 2:
                cell.fill = ZEBRA_FILL
        height = max((_text_height(str(values[c - 1] or ""), 60) for c in wrap), default=24)
        ws.row_dimensions[r].height = min(90, max(24, height))
    return header_row + len(rows) + 2


def _group_values(rows: list[dict] | None) -> list[list[Any]]:
    return [
        [r.get("label"), r.get("count"), r.get("avg_views"), r.get("avg_score")]
        for r in rows or []
    ]


def _write_summary(ws: Worksheet, summary: dict, *, post_run: bool = False, moment: datetime | None = None) -> None:
    for letter, width in zip("ABCD", (34, 14, 18, 24)):
        ws.column_dimensions[letter].width = width
    ws.column_dimensions["C"].hidden = post_run  # bài đăng không có lượt xem
    ws.merge_cells("A1:D1")
    _put(ws, 1, 1, SHEET_SUMMARY).font = Font(name="Calibri", size=20, bold=True, color="234E63")
    ws.row_dimensions[1].height = 36
    ws.merge_cells("A2:D2")
    subtitle = "Bài đăng" if post_run else "Video"
    if moment is not None:
        subtitle += f" • {moment:%d/%m/%Y %H:%M} (giờ Việt Nam)"
    _put(ws, 2, 1, subtitle).font = Font(name="Calibri", size=11, color="607785")
    ws.row_dimensions[2].height = 24
    ws.freeze_panes = "B6"
    headers = ["Số bài" if post_run else "Số video", "View trung bình", "Điểm công thức trung bình"]
    fmts = [None, FMT_INT, FMT_INT, FMT_DEC1]
    row = _write_table(
        ws, 4, "Bảng 1 – Theo nhóm chủ đề (điểm giảm dần)", ["Nhóm chủ đề", *headers],
        _group_values(summary.get("by_topic")), fmts,
    )
    topics = summary.get("by_topic") or []
    if len(topics) >= 2:
        chart = BarChart()
        chart.type = "bar"
        chart.style = 10
        chart.title = "Số bài theo nhóm chủ đề" if post_run else "Số video theo nhóm chủ đề"
        # openpyxl's x axis is the category axis, even for a horizontal bar chart.
        chart.x_axis.title = "Nhóm chủ đề"
        chart.y_axis.title = "Số bài" if post_run else "Số video"
        chart.x_axis.delete = False
        chart.y_axis.delete = False
        chart.x_axis.tickLblPos = "nextTo"
        chart.y_axis.tickLblPos = "nextTo"
        chart.y_axis.numFmt = "0"
        chart.y_axis.majorUnit = max(1, math.ceil(max(r.get("count") or 0 for r in topics) / 6))
        chart.y_axis.scaling.min = 0
        chart.add_data(Reference(ws, min_col=2, min_row=5, max_row=5 + len(topics)), titles_from_data=True)
        chart.set_categories(Reference(ws, min_col=1, min_row=6, max_row=5 + len(topics)))
        chart.series[0].graphicalProperties.solidFill = "55A58B"
        chart.series[0].graphicalProperties.line.noFill = True
        chart.legend = None
        chart.width = 21
        chart.height = max(8, min(16, len(topics) * 0.7 + 4))
        ws.add_chart(chart, "F4")
        ws.column_dimensions["E"].width = 3
        for letter in "FGHIJKLMNO":
            ws.column_dimensions[letter].width = 10
    row = _write_table(
        ws, row, "Bảng 2 – Theo kiểu hook (điểm giảm dần)", ["Kiểu hook", *headers],
        _group_values(summary.get("by_hook")), fmts,
    )
    row = _write_table(
        ws, row, "Bảng 3 – Theo cấu trúc (điểm giảm dần)", ["Cấu trúc", *headers],
        _group_values(summary.get("by_structure")), fmts,
    )

    _put(ws, row, 1, "Nhận xét").font = TITLE_FONT
    text = _text(summary.get("commentary")) or "Chưa có nhận xét."
    note_row = row + 1
    ws.merge_cells(start_row=note_row, start_column=1, end_row=note_row, end_column=4)
    note = _put(ws, note_row, 1, text, align=TOP_WRAP)
    note.font = BODY_FONT
    note.fill = PatternFill("solid", fgColor="E8F2EF")
    note.border = ROW_BORDER
    ws.row_dimensions[note_row].height = min(240, _text_height(text, chars_per_line=80))
    if len(topics) >= 2:
        ws.print_area = f"A1:O{max(35, note_row + 8)}"


def _text_height(text: str, chars_per_line: int) -> float:
    # Merged cells do not auto-fit in Excel, so estimate the height from the text length.
    lines = sum(max(1, math.ceil(len(part) / chars_per_line)) for part in text.split("\n"))
    return 15.0 * lines + 6


def _fmt_time(value: Any) -> str | None:
    if value is None or value == "":
        return None
    if isinstance(value, datetime):
        return value.strftime("%d/%m/%Y %H:%M:%S")
    if isinstance(value, (int, float)):
        return to_vn_datetime(int(value)).strftime("%d/%m/%Y %H:%M:%S")
    return str(value)


def _page_row(page: dict) -> list[Any]:
    return [
        page.get("name") or page.get("url"),
        page.get("url"),
        page.get("followers"),
        page.get("scanned"),
        page.get("passed"),
        page.get("analyzed"),
        page.get("skipped_reason"),
        "\n".join(t for t in (_text(e) for e in (page.get("errors") or [])) if t),
    ]


def _write_run_log(ws: Worksheet, run_log: dict) -> None:
    for letter, width in zip("ABCDEFGH", (28, 60, 14, 12, 14, 14, 30, 60)):
        ws.column_dimensions[letter].width = width

    params = run_log.get("params") or {}
    pages = run_log.get("pages") or []
    usage = run_log.get("usage") or {}
    messages = run_log.get("messages") or []

    row = _write_table(
        ws, 1, "Thời gian", ["Mục", "Giá trị"],
        [
            ["Bắt đầu", _fmt_time(run_log.get("started_at"))],
            ["Kết thúc", _fmt_time(run_log.get("finished_at"))],
        ],
        wrap={2},
    )
    row = _write_table(
        ws, row, "Tham số đã chọn", ["Tham số", "Giá trị"],
        [[_PARAM_LABELS.get(key, key),
          ", ".join(KIND_LABELS.get(kind, kind) for kind in sorted(value)) if key == "formats"
          else ("Đạt cả hai điều kiện" if value else "Đạt một trong hai") if key == "match_all"
          else value]
         for key, value in params.items()],
        wrap={2},
    )
    row = _write_table(
        ws, row, "Kênh",
        ["Kênh", "Link", "Follower", "Đã quét", "Đạt ngưỡng", "Đã phân tích", "Bỏ qua vì", "Lỗi"],
        [_page_row(p) for p in pages],
        [None, None, FMT_INT, FMT_INT, FMT_INT, FMT_INT, None, None],
        wrap={7, 8},
    )
    for r in range(row - len(pages) - 1, row - 1):
        error = ws.cell(r, 8)
        if error.value:
            error.fill = RISK_FILL
            error.font = Font(name="Calibri", color="9C3030", size=11)
    row = _write_table(
        ws, row, "Sử dụng", ["Token Jev", "Token LLM", "Phút Deepgram"],
        [[usage.get("jev_tokens"), usage.get("llm_tokens"), usage.get("deepgram_minutes")]],
        [FMT_INT, FMT_INT, FMT_DEC2],
    )
    _write_table(
        ws, row, "Thông báo", ["STT", "Nội dung"],
        [[i, message] for i, message in enumerate(messages, start=1)],
        wrap={2},
    )
    ws.freeze_panes = "B3"
