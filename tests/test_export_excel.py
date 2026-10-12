"""Offline tests for export_excel.py: sheets, header order, sort, risk fill, formats, filename."""
from __future__ import annotations

import re
from datetime import datetime
from pathlib import Path

import pytest
from openpyxl import load_workbook

import export_excel as xl
from models import Video

NOW = datetime(2026, 10, 12, 9, 5)

HEADERS = [
    "STT", "Kênh", "Follower kênh", "Ngày đăng", "Định dạng", "Thời lượng (giây)", "Link bài",
    "View", "Cảm xúc", "Bình luận", "Chia sẻ", "Tổng tương tác", "View/Follower", "Hệ số so trung vị kênh",
    "Tỷ lệ tương tác", "Tỷ lệ share",
    "Chủ đề bài", "Tóm tắt", "Caption", "Nội dung thoại", "Hook gốc",
    "Góc nhìn", "Mẫu hook điền chỗ trống", "Vị trí CTA",
    "Nhóm chủ đề", "Kiểu hook", "Độ mạnh hook", "Cấu trúc", "CTA",
    "Liên quan ngành (%)", "Làm lại được (%)", "Điểm công thức (0–100)", "Điểm Jev (0–100)", "Cờ rủi ro",
    "Lý do viral", "Gợi ý 1", "Gợi ý 2", "Gợi ý 3",
    "Ghi chú",
]

CAPTION_B = '=HYPERLINK("http://x","bấm") Khỏi hẳn bệnh'

SUMMARY = {
    "by_topic": [
        {"label": "Kiến thức bệnh", "count": 2, "avg_views": 70_000, "avg_score": 62.5},
        {"label": "Pháp lý GPP", "count": 1, "avg_views": None, "avg_score": 90.0},
    ],
    "by_hook": [
        {"label": "Câu hỏi", "count": 2, "avg_views": 70_000, "avg_score": 62.5},
    ],
    "by_structure": [
        {"label": "Liệt kê mẹo", "count": 2, "avg_views": 95_000, "avg_score": 80.0},
    ],
    "commentary": "Công thức câu hỏi đang thắng.",
}

RUN_LOG = {
    "started_at": 1_760_000_000,
    "finished_at": "2026-10-12T08:27:31.344983+00:00",
    "params": {
        "days": 14,
        "formats": {"video", "reel"},
        "min_views": 10000,
        "min_interactions": 0,
    },
    "pages": [
        {"url": "https://fb/p1", "name": "Kênh A", "followers": 200_000, "scanned": 40,
         "passed": 3, "analyzed": 3, "skipped_reason": None, "errors": []},
        {"url": "https://fb/p2", "name": "Kênh C", "followers": 5_000, "scanned": 0,
         "passed": 0, "analyzed": 0, "skipped_reason": "Link không hợp lệ",
         "errors": []},
        {"url": "https://fb/p3", "name": "", "followers": None, "scanned": 2,
         "passed": 0, "analyzed": 0, "skipped_reason": None,
         "errors": ["Không kết nối được tới Facebook"]},
    ],
    "usage": {"jev_tokens": 12_345, "llm_tokens": 6_789, "deepgram_minutes": 3.5, "llm_model": "claude-sonnet"},
    "messages": ["Đã dùng cache cho 2 video"],
}


def _videos() -> list[Video]:
    # Input order is deliberately not the score order, and one video has no score.
    return [
        Video(
            video_id="c", url="https://www.facebook.com/reel/c",
            page_url="https://www.facebook.com/kenhc", page_name="Kênh C",
            page_followers=90_000, posted_at=None, format="video", views=1_000,
            caption="Nội dung không có điểm", score=None,
            jev={},
            llm={"_raw": "Không phải JSON: {oops"},
        ),
        Video(
            video_id="b", url="https://www.facebook.com/reel/b",
            page_url="https://www.facebook.com/kenhb", page_name="Kênh B",
            page_followers=300_000, posted_at=1_760_000_000, format="video",
            duration_s=55.0, views=120_000, reactions=2_500, comments=300, shares=900,
            ratio_follow=0.4, ratio_median=5.0, engagement=0.03, share_rate=0.0075,
            caption=CAPTION_B, transcript="", transcript_status="failed", hook="Khỏi hẳn bệnh",
            score=90.0,
            jev={
                "topic": "Pháp lý GPP (chưa chắc)", "hook_type": "Cảnh báo sai lầm",
                "hook_strength": 1, "structure": "Liệt kê mẹo", "cta": "Bình luận",
                "relevant_p": 0.95, "replicable_p": 0.7,
                "risk_flags": ["Khẳng định tuyệt đối", "So sánh đối thủ"],
                "uncertain": ["topic"],
            },
            notes=["Chưa đủ 5 video để tính trung vị"],
            llm={"chu_de_bai": "Cảnh báo sai lầm", "tom_tat": "Tóm tắt B"},
        ),
        Video(
            video_id="a", url="https://www.facebook.com/reel/a",
            page_url="https://www.facebook.com/kenha", page_name="Kênh A",
            page_followers=200_000, posted_at=1_760_000_000, format="reel",
            duration_s=30.0, views=80_000, reactions=1_000, comments=50, shares=40,
            ratio_follow=0.4, ratio_median=3.2, engagement=0.03, share_rate=0.0005,
            caption="Bạn có biết dấu hiệu này không?", transcript="Thoại A",
            transcript_status="ok", hook="Bạn có biết", score=55.0,
            jev={
                "topic": "Kiến thức bệnh", "hook_type": "Câu hỏi", "hook_strength": 2,
                "structure": "Vấn đề - giải pháp", "cta": "Inbox", "relevant_p": 0.9,
                "replicable_p": 0.8, "risk_flags": [], "uncertain": [],
            },
            llm={
                "chu_de_bai": "Dấu hiệu nhận biết bệnh",
                "tom_tat": "Tóm tắt A",
                "hook_goc": "Bạn có biết…",
                "goc_nhin": "Góc nhìn của dược sĩ tư vấn tại quầy",
                "mau_hook_dien_cho_trong": "Bạn có biết [dấu hiệu] báo hiệu [vấn đề] không?",
                "vi_tri_cta": "Cuối video, sau phần giải thích dấu hiệu",
                "ly_do_viral": "Giả thuyết A",
                "goi_y_nhan_ban": [
                    {"hook_moi": "Chủ nhà thuốc nào cũng…", "y_tuong": "Kể tình huống"},
                    {"hook_moi": "H2", "y_tuong": "Y2"},
                    {"hook_moi": "H3", "y_tuong": "Y3"},
                ],
            },
        ),
        Video(
            video_id="d", url="https://www.facebook.com/reel/d",
            page_url="https://www.facebook.com/kenhd", page_name="Kênh D",
            page_followers=120_000, posted_at=1_760_000_000, format="reel",
            views=70_000, score=70.0, caption="Chia sẻ mẹo nhỏ",
            jev={
                "topic": "Kiến thức bệnh", "hook_type": "Câu hỏi", "hook_strength": 1,
                "structure": "Liệt kê mẹo", "cta": "Lưu", "relevant_p": 0.7,
                "replicable_p": 0.6, "risk_flags": [], "uncertain": [],
            },
        ),
    ]


@pytest.fixture
def exported(tmp_path: Path) -> Path:
    return xl.export(_videos(), SUMMARY, RUN_LOG, out_dir=tmp_path, now=NOW)


@pytest.fixture
def wb(exported: Path):
    return load_workbook(exported)


@pytest.fixture
def detail(wb):
    return wb["Chi tiết video"]


def _col(header: str) -> int:
    return HEADERS.index(header) + 1


def _find_row(ws, text: str, column: int = 1) -> int:
    for r in range(1, ws.max_row + 1):
        if ws.cell(r, column).value == text:
            return r
    raise AssertionError(f"không thấy {text!r} trong cột {column}")


def _width(ws, column: int) -> float | None:
    # openpyxl groups adjacent equal-width columns on save, so look up by range.
    for dim in ws.column_dimensions.values():
        if dim.min is not None and dim.max is not None and dim.min <= column <= dim.max:
            return dim.width
    return None


def test_filename_matches_pattern(exported: Path, tmp_path: Path) -> None:
    assert exported.name == "cong-thuc_2026-10-12_0905.xlsx"
    assert re.fullmatch(r"cong-thuc_\d{4}-\d{2}-\d{2}_\d{4}\.xlsx", exported.name)
    assert exported.parent == tmp_path


def test_existing_file_is_never_overwritten(tmp_path: Path) -> None:
    first = xl.export([], {}, {}, out_dir=tmp_path, now=NOW)
    second = xl.export([], {}, {}, out_dir=tmp_path, now=NOW)
    assert first.name == "cong-thuc_2026-10-12_0905.xlsx"
    assert second.name == "cong-thuc_2026-10-12_0905_2.xlsx"
    assert first.exists() and second.exists()


def test_post_export_labels_and_detail_title(tmp_path):
    post = Video("p1", "https://facebook.test/p1", "page", format="post", interactions=250, jev={"jev_viral": 0.75})
    path = xl.export([post], {}, {"params": {"formats": ["post"], "min_interactions": 200}}, out_dir=tmp_path, now=NOW)
    workbook = load_workbook(path)
    try:
        assert workbook.sheetnames[:2] == ["Tổng quan", "Chi tiết bài đăng"]
        detail = workbook["Chi tiết bài đăng"]
        assert detail.cell(3, _col("Định dạng")).value == "Bài đăng"
        assert detail.cell(3, _col("Tổng tương tác")).value == 250
        assert detail.cell(3, _col("Điểm Jev (0–100)")).value == 75
        assert workbook["Tổng hợp công thức"].cell(5, 2).value == "Số bài"
        for header in xl._VIDEO_ONLY_HEADERS:
            assert detail.column_dimensions[xl.get_column_letter(_col(header))].hidden
    finally:
        workbook.close()


def test_run_log_shows_ratio_and_combination(tmp_path):
    video = Video("v1", "https://facebook.test/v1", "page", views=50_000, page_followers=1_000)
    params = {"formats": ["video"], "min_views": 50_000, "min_ratio_follow": 5, "match_all": False}
    path = xl.export([video], {}, {"params": params}, out_dir=tmp_path, now=NOW)
    workbook = load_workbook(path)
    try:
        cells = [cell for sheet in workbook for row in sheet.iter_rows(values_only=True) for cell in row if cell is not None]
    finally:
        workbook.close()
    assert "Lượt xem tối thiểu so với follower (×)" in cells and 5 in cells
    assert "Đạt một trong hai" in cells


def test_sheet_names_in_order(wb) -> None:
    assert wb.sheetnames == ["Tổng quan", "Chi tiết video", "Tổng hợp công thức", "Nhật ký"]
    assert wb.active.title == "Tổng quan"


def test_detail_header_order_matches_plan(detail) -> None:
    assert detail.max_column == len(HEADERS) == 39
    assert [detail.cell(2, c).value for c in range(1, len(HEADERS) + 1)] == HEADERS


def test_rows_sorted_by_score_desc_with_missing_score_last(detail) -> None:
    assert detail.max_row == 6
    assert [detail.cell(r, _col("STT")).value for r in range(3, 7)] == [1, 2, 3, 4]
    assert [detail.cell(r, _col("Kênh")).value for r in range(3, 7)] == [
        "Kênh B", "Kênh D", "Kênh A", "Kênh C",
    ]
    assert [detail.cell(r, _col("Điểm công thức (0–100)")).value for r in range(3, 7)] == [
        90.0, 70.0, 55.0, None,
    ]


def test_freeze_panes_and_autofilter(detail) -> None:
    assert detail.freeze_panes == "C3"
    assert detail.auto_filter.ref == "A2:AM6"


def test_risk_highlight_is_limited_to_nonempty_flags(detail) -> None:
    risky = detail.cell(3, _col("Cờ rủi ro"))
    assert risky.fill.fill_type == "solid"
    assert risky.fill != detail.cell(3, _col("Kênh")).fill
    assert detail.cell(4, _col("Cờ rủi ro")).value is None


def test_data_rows_have_readable_borders_and_capped_height(detail) -> None:
    for row in range(3, 7):
        assert detail.cell(row, _col("Kênh")).border.bottom.style
        assert 24 <= detail.row_dimensions[row].height <= 90


def test_number_and_percent_formats(detail) -> None:
    row = 3  # Kênh B: has every metric filled in
    assert detail.cell(row, _col("View")).number_format == "#,##0"
    assert detail.cell(row, _col("Follower kênh")).number_format == "#,##0"
    assert detail.cell(row, _col("Thời lượng (giây)")).number_format == "#,##0"
    assert detail.cell(row, _col("View/Follower")).number_format == "0.0%"
    assert detail.cell(row, _col("View/Follower")).value == pytest.approx(0.4)
    assert detail.cell(row, _col("Hệ số so trung vị kênh")).number_format == "#,##0.00"
    assert detail.cell(row, _col("Tỷ lệ tương tác")).number_format == "0.0%"
    assert detail.cell(row, _col("Tỷ lệ share")).number_format == "0.00%"
    assert detail.cell(row, _col("Tỷ lệ share")).value == pytest.approx(0.0075)
    assert detail.cell(row, _col("Liên quan ngành (%)")).number_format == "0%"
    assert detail.cell(row, _col("Liên quan ngành (%)")).value == pytest.approx(0.95)
    assert detail.cell(row, _col("Làm lại được (%)")).number_format == "0%"
    assert detail.cell(row, _col("Điểm công thức (0–100)")).number_format == "#,##0.0"


def test_long_text_columns_wrap_at_width_60(detail) -> None:
    for header in (
        "Tóm tắt", "Caption", "Nội dung thoại", "Hook gốc",
        "Góc nhìn", "Mẫu hook điền chỗ trống", "Vị trí CTA", "Gợi ý 1", "Ghi chú",
    ):
        col = _col(header)
        assert _width(detail, col) == 60, header
        assert detail.cell(3, col).alignment.wrap_text is True, header


def test_short_columns_keep_default_width_without_wrap(detail) -> None:
    col = _col("Kênh")
    assert _width(detail, col) == 24
    assert not detail.cell(3, col).alignment.wrap_text


def test_formula_like_caption_stays_literal_text(detail) -> None:
    assert detail.cell(3, _col("Caption")).value == CAPTION_B


def test_detail_values_map_from_jev_llm_and_metrics(detail) -> None:
    row = 5  # Kênh A
    assert detail.cell(row, _col("Ngày đăng")).value == datetime(2025, 10, 9, 15, 53, 20)
    assert detail.cell(row, _col("Định dạng")).value == "Reel"
    assert detail.cell(row, _col("Chủ đề bài")).value == "Dấu hiệu nhận biết bệnh"
    assert detail.cell(row, _col("Hook gốc")).value == "Bạn có biết…"
    assert detail.cell(row, _col("Độ mạnh hook")).value == 2
    assert detail.cell(row, _col("Nhóm chủ đề")).value == "Kiến thức bệnh"
    assert detail.cell(row, _col("CTA")).value == "Inbox"
    assert detail.cell(row, _col("Lý do viral")).value == "Giả thuyết A"
    assert detail.cell(row, _col("Gợi ý 1")).value == (
        "Hook: Chủ nhà thuốc nào cũng…\nÝ tưởng: Kể tình huống"
    )
    assert detail.cell(row, _col("Gợi ý 3")).value == "Hook: H3\nÝ tưởng: Y3"
    assert detail.cell(row, _col("Ghi chú")).value is None


@pytest.mark.parametrize(("header", "expected"), [
    ("Góc nhìn", "Góc nhìn của dược sĩ tư vấn tại quầy"),
    ("Mẫu hook điền chỗ trống", "Bạn có biết [dấu hiệu] báo hiệu [vấn đề] không?"),
    ("Vị trí CTA", "Cuối video, sau phần giải thích dấu hiệu"),
])
def test_analysis_formula_fields_export_from_llm(detail, header, expected) -> None:
    assert detail.cell(5, _col(header)).value == expected
    assert detail.cell(6, _col(header)).value is None


def test_risk_flags_uncertain_labels_and_failed_transcript_in_notes(detail) -> None:
    row = 3  # Kênh B
    assert detail.cell(row, _col("Cờ rủi ro")).value == "Khẳng định tuyệt đối, So sánh đối thủ"
    assert detail.cell(row, _col("Nhóm chủ đề")).value == "Pháp lý GPP (chưa chắc)"
    notes = detail.cell(row, _col("Ghi chú")).value
    assert notes.splitlines() == [
        "Chưa đủ 5 video để tính trung vị",
        "Nhãn chưa chắc: Nhóm chủ đề",
        "Chưa lấy được thoại",
    ]


def test_missing_jev_llm_fields_are_blank_and_raw_goes_to_notes(detail) -> None:
    row = 6  # Kênh C: no score, empty jev, raw LLM text only
    assert detail.cell(row, _col("Điểm công thức (0–100)")).value is None
    assert detail.cell(row, _col("Chủ đề bài")).value is None
    assert detail.cell(row, _col("Gợi ý 1")).value is None
    assert detail.cell(row, _col("Ngày đăng")).value is None
    notes = detail.cell(row, _col("Ghi chú")).value
    assert "Không phải JSON: {oops" in notes


def test_empty_export_still_writes_four_sheets(tmp_path: Path) -> None:
    path = xl.export([], {}, {}, out_dir=tmp_path, now=NOW)
    wb = load_workbook(path)
    assert wb.sheetnames == ["Tổng quan", "Chi tiết video", "Tổng hợp công thức", "Nhật ký"]
    assert wb["Chi tiết video"].max_row == 2


def test_summary_sheet_tables_and_commentary(wb) -> None:
    ws = wb["Tổng hợp công thức"]
    topic_title = _find_row(ws, "Bảng 1 – Theo nhóm chủ đề (điểm giảm dần)")
    assert [ws.cell(topic_title + 1, c).value for c in range(1, 5)] == [
        "Nhóm chủ đề", "Số video", "View trung bình", "Điểm công thức trung bình",
    ]
    assert [ws.cell(topic_title + 2, c).value for c in range(1, 5)] == [
        "Kiến thức bệnh", 2, 70_000, 62.5,
    ]
    assert ws.cell(topic_title + 2, 3).number_format == "#,##0"
    assert ws.cell(topic_title + 2, 4).number_format == "#,##0.0"
    assert [ws.cell(topic_title + 3, c).value for c in range(1, 5)] == [
        "Pháp lý GPP", 1, None, 90.0,
    ]

    hook_title = _find_row(ws, "Bảng 2 – Theo kiểu hook (điểm giảm dần)")
    assert ws.cell(hook_title + 2, 1).value == "Câu hỏi"
    structure_title = _find_row(ws, "Bảng 3 – Theo cấu trúc (điểm giảm dần)")
    assert [ws.cell(structure_title + 2, c).value for c in range(1, 5)] == [
        "Liệt kê mẹo", 2, 95_000, 80.0,
    ]
    assert ws.cell(structure_title + 2, 4).number_format == "#,##0.0"
    assert ws.cell(structure_title + 2, 1).fill.fgColor == ws.cell(hook_title + 2, 1).fill.fgColor
    assert ws.cell(structure_title + 2, 1).fill.patternType == "solid"

    note = _find_row(ws, "Nhận xét")
    assert ws.cell(note + 1, 1).value == "Công thức câu hỏi đang thắng."


def test_run_log_sheet_contents(wb) -> None:
    ws = wb["Nhật ký"]
    assert ws.cell(_find_row(ws, "Kết thúc"), 2).value == datetime(2026, 10, 12, 15, 27, 31, 345000)
    assert ws.cell(_find_row(ws, "Bắt đầu"), 2).value == datetime(2025, 10, 9, 15, 53, 20)
    assert ws.cell(_find_row(ws, "Bắt đầu"), 2).number_format == xl.FMT_DATETIME

    assert ws.cell(_find_row(ws, "Lượt xem tối thiểu"), 2).value == 10000
    assert ws.cell(_find_row(ws, "Tương tác tối thiểu"), 2).value == 0

    skipped = _find_row(ws, "Kênh C")
    assert ws.cell(skipped, 7).value == "Link không hợp lệ"
    failed = _find_row(ws, "https://fb/p3")
    assert ws.cell(failed, 8).value == "Không kết nối được tới Facebook"

    usage_title = _find_row(ws, "Sử dụng")
    assert [ws.cell(usage_title + 1, c).value for c in range(1, 5)] == [
        "Token Jev", "Token LLM", "Phút Deepgram", "Model LLM",
    ]
    assert [ws.cell(usage_title + 2, c).value for c in range(1, 5)] == [12_345, 6_789, 3.5, "claude-sonnet"]
    assert ws.cell(usage_title + 2, 1).number_format == "#,##0"
    assert ws.cell(usage_title + 2, 3).number_format == "#,##0.00"

    message_row = _find_row(ws, "Đã dùng cache cho 2 video", column=2)
    assert ws.cell(message_row, 2).value == "Đã dùng cache cho 2 video"


def test_build_summary_tables_groups_and_averages() -> None:
    videos = [
        Video(video_id="1", url="", page_url="", views=100, score=60.0,
              jev={"topic": "Kiến thức bệnh (chưa chắc)", "hook_type": "Câu hỏi", "structure": "Liệt kê (chưa chắc)"}),
        Video(video_id="2", url="", page_url="", views=None, score=80.0,
              jev={"topic": "Kiến thức bệnh", "hook_type": "Câu hỏi", "structure": "Liệt kê"}),
        Video(video_id="3", url="", page_url="", views=50, score=None, jev={}),
    ]
    tables = xl.build_summary_tables(videos)
    assert tables == {
        "by_topic": [
            {"label": "Kiến thức bệnh", "count": 2, "avg_views": 100.0, "avg_score": 70.0},
            {"label": "Chưa gắn nhãn", "count": 1, "avg_views": 50.0, "avg_score": None},
        ],
        "by_hook": [
            {"label": "Câu hỏi", "count": 2, "avg_views": 100.0, "avg_score": 70.0},
            {"label": "Chưa gắn nhãn", "count": 1, "avg_views": 50.0, "avg_score": None},
        ],
        "by_structure": [
            {"label": "Liệt kê", "count": 2, "avg_views": 100.0, "avg_score": 70.0},
            {"label": "Chưa gắn nhãn", "count": 1, "avg_views": 50.0, "avg_score": None},
        ],
        "commentary": "",
    }


def test_build_summary_tables_empty() -> None:
    assert xl.build_summary_tables([]) == {"by_topic": [], "by_hook": [], "by_structure": [], "commentary": ""}


@pytest.mark.parametrize("field,table", [("topic", "by_topic"), ("hook_type", "by_hook"), ("structure", "by_structure")])
def test_summary_ranks_score_then_count_with_missing_last(field, table):
    videos = [
        Video(str(i), "", "", score=score, jev={field: label})
        for i, (label, score) in enumerate([
            ("Popular", 20), ("Popular", 20), ("Popular", 20),
            ("Best", 90), ("Best", 90), ("Tie", 90), ("Missing", None), ("Zero", 0),
        ])
    ]
    rows = xl.build_summary_tables(videos)[table]
    assert [row["label"] for row in rows] == ["Best", "Tie", "Popular", "Zero", "Missing"]

def test_live_labels_in_details_and_log_params(tmp_path):
    videos = [Video(kind, kind, "page", format=kind.removesuffix("_live"), is_live=kind.endswith("_live"))
              for kind in ("video", "reel", "video_live", "reel_live", "post")]
    path = xl.export(videos, {}, {"params": {"formats": ["video_live", "reel_live"]}}, out_dir=tmp_path, now=NOW)
    workbook = load_workbook(path)
    try:
        detail = workbook[xl.SHEET_DETAIL]
        assert {detail.cell(row, _col("Định dạng")).value for row in range(3, 8)} == {
            "Video", "Reel", "Video (Live)", "Reel (Live)", "Bài đăng"}
        values = [cell.value for row in workbook[xl.SHEET_LOG] for cell in row]
        assert "Reel (Live), Video (Live)" in values
    finally:
        workbook.close()


def test_overview_kpis_charts_and_print_layout(wb):
    ws = wb[xl.SHEET_OVERVIEW]
    expected = {
        "Số kênh": 3, "Số video quét được": 42, "Đạt ngưỡng": 3,
        "Đã phân tích": 3, "Số mục xuất": 4, "Gợi ý giữ lại": 3, "Video/bài có gợi ý": 1,
        "Token Jev": 12_345, "Token LLM": 6_789, "Phút Deepgram": 3.5,
        "Model LLM": "claude-sonnet", "Số ngày quét": 14, "Lượt xem tối thiểu": 10_000,
    }
    for label, value in expected.items():
        assert ws.cell(_find_row(ws, label), 2).value == value
    detail = wb[xl.SHEET_DETAIL]
    exported_counts = [
        sum(detail.cell(row, _col(f"Gợi ý {index}")).value is not None for index in range(1, 4))
        for row in range(3, detail.max_row + 1)
    ]
    assert ws.cell(_find_row(ws, "Gợi ý giữ lại"), 2).value == sum(exported_counts)
    assert ws.cell(_find_row(ws, "Video/bài có gợi ý"), 2).value == sum(count > 0 for count in exported_counts)
    assert not any("điểm ≥ 70" in str(cell.value) for row in ws for cell in row)
    assert ws.cell(_find_row(ws, "Phút Deepgram"), 2).number_format == xl.FMT_DEC2
    assert len(ws._charts) >= 1
    assert ws._charts[0].series[0].val.numRef.f.startswith("'Tổng quan'!")
    assert ws.sheet_view.showGridLines is False
    for sheet in wb:
        assert sheet.freeze_panes
        assert sheet.print_title_rows
        assert sheet.page_setup.orientation == "landscape"
        assert sheet.page_setup.fitToWidth == 1
        assert sheet.sheet_properties.tabColor is not None


def test_detail_links_and_conditional_formats_preserved(detail):
    for row in range(3, 7):
        cell = detail.cell(row, _col("Link bài"))
        assert cell.hyperlink.target == cell.value
        assert cell.hyperlink.display == "Mở bài ↗"
    assert any(rule.type == "dataBar" for rules in detail.conditional_formatting._cf_rules.values() for rule in rules)
    assert any(rule.type == "expression" for rules in detail.conditional_formatting._cf_rules.values() for rule in rules)


def test_read_summary_round_trip(exported):
    result = xl.read_summary(exported)
    assert result["kind"] == "video"
    assert result["count"] == 4
    assert result["created"] == NOW
    assert result["usage"] == RUN_LOG["usage"]
    expected = sorted(_videos(), key=lambda v: (v.score is None, -(v.score or 0)))
    assert result["items"] == [
        {"title": v.llm.get("chu_de_bai") or v.caption or v.url,
         "page": v.page_name or v.page_url, "url": v.url, "views": v.views,
         "interactions": v.interactions, "score": v.score, "topic": v.jev.get("topic", "")}
        for v in expected
    ]


def test_read_summary_post_and_empty_exports(tmp_path):
    post = Video("p", "https://facebook.test/p", "page", format="post", caption="Bài mới",
                 interactions=250, score=80, jev={"topic": "Mẹo"}, llm={"chu_de_bai": "Mẹo hay"})
    path = xl.export([post], {}, {"params": {"formats": ["post"]}}, out_dir=tmp_path, now=NOW)
    result = xl.read_summary(path)
    assert result["kind"] == "post"
    assert result["count"] == 1
    assert result["items"][0] == {
        "title": "Mẹo hay", "page": "page", "url": post.url, "views": None,
        "interactions": 250, "score": 80.0, "topic": "Mẹo",
    }
    assert result["usage"] == {"jev_tokens": 0, "llm_tokens": 0, "deepgram_minutes": 0.0, "llm_model": ""}
    empty = xl.export([], {}, {}, out_dir=tmp_path, now=NOW)
    assert xl.read_summary(empty)["count"] == 0
    assert xl.read_summary(empty)["items"] == []


def test_read_summary_rejects_garbage_foreign_and_missing_files(tmp_path):
    garbage = tmp_path / "hong.xlsx"
    garbage.write_bytes(b"not an Excel workbook")
    assert xl.read_summary(garbage) == {}
    assert xl.read_summary(tmp_path / "missing.xlsx") == {}
    foreign = tmp_path / "khac.xlsx"
    from openpyxl import Workbook
    workbook = Workbook()
    workbook.save(foreign)
    workbook.close()
    assert xl.read_summary(foreign) == {}


def test_read_summary_rejects_imitation_without_detail_headers(tmp_path):
    from openpyxl import Workbook
    workbook = Workbook()
    workbook.active.title = xl.SHEET_DETAIL
    workbook.create_sheet(xl.SHEET_SUMMARY)
    workbook.create_sheet(xl.SHEET_LOG)
    path = tmp_path / "khong-dung-mau.xlsx"
    workbook.save(path)
    workbook.close()
    assert xl.read_summary(path) == {}
