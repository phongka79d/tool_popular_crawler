import base64
import io
from types import SimpleNamespace

import pytest
from PIL import Image

import db
import jev_client
import llm_client
import export_excel
from collectors import graphql
from models import Video


def photo(url):
    return {"__typename": "Photo", "photo_image": {"uri": url}}


def test_single_photo_ignores_profile_and_video_thumbnail():
    story = {"__typename": "Story", "post_id": "p", "message": {"text": "Caption"},
             "profile_picture": {"uri": "profile"}, "attachments": [{"styles": {"attachment": {
                 "media": {**photo("image"), "preferred_thumbnail": {"uri": "thumb"}}
             }}}]}
    assert graphql.flatten(story)["image_urls"] == ["image"]


def test_album_urls_order_deduplicated_and_capped():
    nodes = [{"media": {"__typename": "Photo", "image": {"uri": str(i)},
                        "viewer_image": {"uri": "fallback"}}} for i in [1, 1, 2, 3, 4, 5]]
    story = {"__typename": "Story", "post_id": "p", "message": {"text": "Caption"},
             "attachments": [{"styles": {"attachment": {"all_subattachments": {"nodes": nodes}}}}]}
    assert graphql.flatten(story)["image_urls"] == ["1", "2", "3", "4"]


def test_picture_only_post_is_kept_but_empty_post_is_dropped():
    story = {"__typename": "Story", "post_id": "p",
             "attachments": [{"styles": {"attachment": {"media": photo("only")}}}]}
    record = graphql.flatten(story)
    assert (record["caption"], record["image_urls"]) == ("", ["only"])
    assert graphql.flatten({"__typename": "Story", "post_id": "q"}) is None


def test_mixed_attachments_preserve_post_image_order():
    story = {"__typename": "Story", "post_id": "p", "message": {"text": "Caption"},
             "attachments": [
                 {"styles": {"attachment": {"all_subattachments": {"nodes": [{"media": photo("first")}]}}}},
                 {"styles": {"attachment": {"media": photo("second")}}},
             ]}
    assert graphql.flatten(story)["image_urls"] == ["first", "second"]



@pytest.mark.parametrize("answer,expected", [("1234", True), ("9999", False)])
def test_vision_checks_number(monkeypatch, answer, expected):
    monkeypatch.setattr(llm_client.secrets, "randbelow", lambda _: 234)
    def create(**kwargs):
        parts = kwargs["messages"][0]["content"]
        assert "1234" not in parts[0]["text"]
        Image.open(io.BytesIO(base64.b64decode(parts[1]["image_url"]["url"].split(",")[1]))).verify()
        return SimpleNamespace(choices=[SimpleNamespace(message=SimpleNamespace(content=answer))], usage=None)
    client = SimpleNamespace(chat=SimpleNamespace(completions=SimpleNamespace(create=create)))
    monkeypatch.setattr(llm_client, "make_client", lambda *args: client)
    assert llm_client.check_vision("endpoint", "key", "model") is expected


def test_vision_image_rejection_returns_false(monkeypatch):
    class Rejection(Exception):
        status_code = 400
    def create(**kwargs):
        raise Rejection("unsupported image content")
    client = SimpleNamespace(chat=SimpleNamespace(completions=SimpleNamespace(create=create)))
    monkeypatch.setattr(llm_client, "make_client", lambda *args: client)
    assert llm_client.check_vision("endpoint", "key", "model") is False


def test_image_cache_separate_and_state_contains_image_text():
    post = Video("p", "u", "page", format="post", caption="Caption")
    text_key = db.analysis_cache_key(post)
    post.transcript, post.transcript_status = "Image words", "images"
    assert db.analysis_cache_key(post) != text_key
    assert jev_client.build_state(post, "page")["transcript"] == "Image words"
    assert "Image words" in llm_client._user_prompt(post, {})
    first = db.analysis_cache_key(post)
    post.transcript = "Changed image words"
    assert db.analysis_cache_key(post) != first


def test_capability_cache_matches_endpoint_and_model():
    llm = {"base_url": "a", "model": "m", "vision": {"base_url": "a", "model": "m", "ok": True}}
    assert llm_client.vision_cached(llm) is True
    llm["model"] = "other"
    assert llm_client.vision_cached(llm) is None


def test_post_export_shows_image_content(tmp_path):
    from openpyxl import load_workbook
    post = Video("p", "u", "page", format="post", transcript="Image words", transcript_status="images")
    path = export_excel.export([post], {}, {"params": {"formats": ["post"]}}, out_dir=tmp_path)
    book = load_workbook(path)
    try:
        sheet = book["Chi tiết bài đăng"]
        cell = next(c for c in sheet[2] if c.value == "Nội dung ảnh")
        assert sheet.cell(3, cell.column).value == "Image words"
        assert not sheet.column_dimensions[cell.column_letter].hidden
    finally:
        book.close()


def test_describe_images_sends_all_resized_images_once():
    data = io.BytesIO()
    Image.new("RGBA", (2048, 512), "white").save(data, format="PNG")
    calls = []
    def create(**kwargs):
        calls.append(kwargs)
        parts = kwargs["messages"][0]["content"]
        assert len(parts) == 3
        for part in parts[1:]:
            image = Image.open(io.BytesIO(base64.b64decode(part["image_url"]["url"].split(",")[1])))
            assert image.size == (1024, 256)
            assert image.format == "JPEG"
        return SimpleNamespace(choices=[SimpleNamespace(message=SimpleNamespace(content="Read image text"))],
                               usage=SimpleNamespace(total_tokens=1800))
    client = SimpleNamespace(chat=SimpleNamespace(completions=SimpleNamespace(create=create)))
    assert llm_client.describe_images(client, "model", [data.getvalue()] * 2, "Caption", "low") == ("Read image text", 1800)
    assert len(calls) == 1
